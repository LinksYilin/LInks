# -*- coding: utf-8 -*-
"""fill_zenodo_doi.py — 把 Zenodo DOI 填入论文与所有投稿材料"""
import os
import re

from docx import Document

ROOT = r'D:\GED_mutation'
DOC = os.path.join(ROOT, 'manuscript_routeA.docx')

CONCEPT = '10.5281/zenodo.23086956'      # 概念 DOI：始终指向最新版
VERSION = '10.5281/zenodo.23086957'      # 本版 DOI
URL = 'https://zenodo.org/records/23086957'
GITHUB = 'https://github.com/LinksYilin/LInks'

# ---------------- 论文 ----------------
DOC_RULES = [
    ('and archived at Zenodo under a DOI to be assigned at acceptance.',
     f'and archived at Zenodo under DOI {CONCEPT} ({URL}).'),
    ('and archived at Zenodo under a DOI to be assigned at acceptance;',
     f'and archived at Zenodo under DOI {CONCEPT} ({URL});'),
    ('released under an MIT licence at https://github.com/LinksYilin/LInks and archived at '
     'Zenodo under a DOI to be assigned at acceptance;',
     f'released under an MIT licence at {GITHUB} and archived at Zenodo under DOI {CONCEPT} '
     f'({URL});'),
]

# ---------------- Markdown ----------------
MD_RULES = [
    ('archived at Zenodo under a DOI to be assigned at acceptance',
     f'archived at Zenodo under DOI {CONCEPT}'),
    ('archived at Zenodo under a DOI to be assigned at acceptance.',
     f'archived at Zenodo under DOI {CONCEPT}.'),
    ('**DOI** (Digital Object Identifier)', '**DOI** (Digital Object Identifier)'),
    ('| Repository: *to be supplied at submission* |', f'| Repository: {GITHUB} |'),
    ('| Archive DOI | *to be supplied at submission* |',
     f'| Archive DOI | {CONCEPT} (concept DOI; version {VERSION}) |'),
    ('- Repository: *to be supplied at submission*', f'- Repository: {GITHUB}'),
    ('- Archive DOI: *to be supplied at submission*',
     f'- Archive DOI: {CONCEPT} (concept DOI, always resolves to the latest version)'),
    ('*to be supplied at submission*', f'{CONCEPT}'),
    ('archived at Zenodo (DOI and URL to be inserted at submission)',
     f'archived at Zenodo (DOI {CONCEPT})'),
    ('archived at Zenodo under a DOI to be assigned at acceptance;',
     f'archived at Zenodo under DOI {CONCEPT};'),
    ('repository (MIT licence), archived at Zenodo',
     f'{GITHUB} (MIT licence), archived at Zenodo'),
    ('doi: ""', f'doi: "{CONCEPT}"'),
    ('"10.xxxx/xxxxx"', f'"{CONCEPT}"'),
]


def patch_docx():
    d = Document(DOC)
    n = 0
    for p in d.paragraphs:
        if not p.runs:
            continue
        full = ''.join(r.text for r in p.runs)
        new = full
        for a, b in DOC_RULES:
            if a in new:
                new = new.replace(a, b)
                n += 1
        if new != full:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ''
    d.save(DOC)
    print(f'论文: {n} 处')
    return n


def patch_md():
    files = ['补充材料_Supplementary.md', '标题页与投稿清单_TitlePage_Checklist.md',
             'CITATION.cff', '投稿信_CoverLetter.md', 'README.md',
             '如何拿到DOI_三步指南.md', '发布与归档操作指南.md',
             'Zenodo手动上传材料.md']
    total = 0
    for f in files:
        p = os.path.join(ROOT, f)
        if not os.path.exists(p):
            continue
        t = open(p, encoding='utf-8').read()
        orig = t
        c = 0
        for a, b in MD_RULES:
            if a in t and a != b:
                t = t.replace(a, b)
                c += 1
        if t != orig:
            open(p, 'w', encoding='utf-8').write(t)
            total += c
            print(f'  {f}: {c} 处')
    return total


print('=== 填入 DOI ===')
patch_docx()
patch_md()

print()
print('=== 校验 ===')
d = Document(DOC)
txt = '\n'.join(p.text for p in d.paragraphs)
print(f'论文含概念 DOI: {txt.count(CONCEPT)} 处')
for f in ['补充材料_Supplementary.md', '标题页与投稿清单_TitlePage_Checklist.md', 'CITATION.cff']:
    t = open(os.path.join(ROOT, f), encoding='utf-8').read()
    print(f'  {f}: {t.count(CONCEPT)} 处')
print()
print(f'概念 DOI: {CONCEPT}')
print(f'版本 DOI: {VERSION}')
