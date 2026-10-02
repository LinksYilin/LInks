# -*- coding: utf-8 -*-
"""audit_submission_figures.py — 核对投稿包的独立图文件与手稿内嵌图是否一致"""
import hashlib
import os
import zipfile

from docx import Document
from docx.oxml.ns import qn

SUB = r'D:\GED_mutation\submission'
DOCX = os.path.join(SUB, '01_manuscript.docx')
FIGDIR = os.path.join(SUB, 'figures')
PUB = r'D:\GED_mutation\figures\publication'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'

d = Document(DOCX)

# 内嵌图 -> 对应图注
embedded = []
for i, p in enumerate(d.paragraphs):
    if p._p.findall('.//' + qn('a:blip')):
        rid = p._p.findall('.//' + qn('a:blip'))[0].get(R + 'embed')
        sha = hashlib.sha256(d.part.rels[rid].target_part.blob).hexdigest()
        # 找后面的图注
        cap = ''
        for j in range(i + 1, min(i + 4, len(d.paragraphs))):
            t = d.paragraphs[j].text.strip()
            if t.startswith('Figure '):
                cap = t.split('|')[0].strip()
                break
        embedded.append((cap, sha))

print('=== 手稿内嵌图 ===')
for cap, sha in embedded:
    print(f'  {cap:<12} sha={sha[:16]}')

print()
print('=== 投稿包独立图文件 ===')
subfigs = sorted(os.listdir(FIGDIR)) if os.path.isdir(FIGDIR) else []
sub_sha = {}
for f in subfigs:
    h = hashlib.sha256(open(os.path.join(FIGDIR, f), 'rb').read()).hexdigest()
    sub_sha[h] = f
    print(f'  {f:<14} sha={h[:16]}')

print()
print('=== 一致性 ===')
ok = 0
for cap, sha in embedded:
    if sha in sub_sha:
        print(f'  ✅ {cap} -> {sub_sha[sha]}')
        ok += 1
    else:
        print(f'  ❌ {cap} 在投稿包中找不到匹配文件')
print(f'\n  匹配 {ok}/{len(embedded)}')

# 反向：投稿包里有但手稿没有的
extra = [f for h, f in sub_sha.items() if h not in {s for _, s in embedded}]
if extra:
    print(f'  ⚠ 投稿包中多出的文件: {extra}')

# 与论文原始图对比
print()
print('=== 与论文原始 PNG 对比 ===')
for f in subfigs:
    n = f.replace('.png', '')
    idx = n.replace('Figure', '')
    cand = [x for x in os.listdir(PUB) if x.startswith('figure' + idx)
            and x.endswith('.png') and 'grayscale' not in x
            and 'limited_increment' not in x and 'model_comparison' not in x]
    if cand:
        a = hashlib.sha256(open(os.path.join(FIGDIR, f), 'rb').read()).hexdigest()
        b = hashlib.sha256(open(os.path.join(PUB, cand[0]), 'rb').read()).hexdigest()
        print(f'  {f:<14} vs {cand[0]:<40} {"✅" if a == b else "❌ 不同"}')
