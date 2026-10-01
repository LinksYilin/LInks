# -*- coding: utf-8 -*-
"""
fix_reviewer2.py — 按审稿人 2 的确认意见修正稿件
==================================================
R2-m5  EGNN 种子数在 Table 1 脚注（2 种子）与正文/补充（3 种子）不一致
R2-M6  摘要未给出形成接触的可重复性（65.7%）
R2-M3  "information budget 模板" 的承诺过强
R2-m2  缺少 Conclusions 节
"""
import os
from docx import Document

DOC = r'D:\GED_mutation\manuscript_routeA.docx'

FIX = [
    # --- R2-m5 ---
    ('\u2021 The equivariant encoder collapsed on a subset of seeds in every definition tested '
     '(for side-chain-centroid graphs, S669 r = 0.393 and \u22120.081 for two seeds; ssym r = 0.068 '
     'and 0.402); its values reflect optimisation instability rather than a representation effect.',
     '\u2021 The equivariant encoder collapsed on a subset of seeds in every definition tested. On '
     'side-chain-centroid graphs its S669 seeds gave r = 0.393, \u22120.081 and 0.079 (mean 0.131, '
     'SD 0.197) and its ssym seeds gave 0.068 and 0.402; a variant whose coordinate-update step was '
     'rescaled from zero at initialisation did not reduce the spread (SD 0.216). Its values reflect '
     'optimisation instability rather than a representation effect.'),

    # --- R2-M6 ---
    ('whereas C\u03b2, side-chain-centroid and all-atom graphs changed in 9.9%, 90.7% and 81.6% of '
     'pairs.',
     'whereas C\u03b2, side-chain-centroid and all-atom graphs changed in 9.9%, 90.7% and 81.6% of '
     'pairs. Across repeated FoldX runs, broken-contact sets agreed at 96.2% and formed-contact sets '
     'at 65.7%.'),

    # --- R2-M3 ---
    ('We provide the audit protocol and the resulting information budget as a template for deciding '
     'where representation-layer effort is likely to pay off.',
     'We provide the audit protocol and the resulting feature-block accounting as a reusable '
     'procedure for weighing representation-layer effort against sequence and physicochemical '
     'alternatives in a new setting.'),
]

CONCLUSIONS = (
    'Conclusions. The atom-level definition of a residue contact graph determines what structural '
    'change a model describes far more strongly than it determines what the model predicts. Within '
    'the encoder family, training data and evaluation protocol tested here, and using contact graphs '
    'built from FoldX-modelled mutants on fixed backbones, twelve of the thirteen paired definition '
    'comparisons were non-significant, and adding a contact-graph branch to an ESM-2 650M sequence '
    'baseline produced no increment larger than |\u0394r| \u2248 0.04 on S669 and \u2248 0.07 on '
    'ssym. The single structural quantity that correlated with \u0394\u0394G on its own, the count '
    'of broken hydrophobic contacts, carried no information beyond five physicochemical features. '
    'We did not evaluate any published structure-based predictor or any structural stream other than '
    'the one FoldX produced, so these bounds apply to the representation layer as tested rather than '
    'to structural information in general. The audit protocol developed here, combining a capacity '
    'ladder with paired equivalence testing and feature-block accounting, transfers to other '
    'representation choices in mutation-effect modelling.'
)


def main():
    d = Document(DOC)
    n = 0
    for p in d.paragraphs:
        if not p.runs:
            continue
        full = ''.join(r.text for r in p.runs)
        new = full
        for a, b in FIX:
            if a in new:
                new = new.replace(a, b)
                n += 1
        if new != full:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ''

    # --- R2-m2：在 Limitations 之后插入 Conclusions ---
    paras = list(d.paragraphs)
    anchor = None
    for p in paras:
        if p.text.strip().startswith('Data & Code Availability'):
            anchor = p
            break
    if anchor is not None:
        from docx.oxml import OxmlElement
        from docx.text.paragraph import Paragraph
        newp = OxmlElement('w:p')
        anchor._p.addprevious(newp)
        np_ = Paragraph(newp, anchor._parent)
        try:
            np_.style = d.paragraphs[0].style
        except Exception:
            pass
        np_.add_run(CONCLUSIONS)
        n += 1
        print('  Conclusions 节已插入')
    else:
        print('  ⚠ 未找到 Data & Code Availability 锚点')

    d.save(DOC)
    print(f'共修正 {n} 处')


if __name__ == '__main__':
    main()
