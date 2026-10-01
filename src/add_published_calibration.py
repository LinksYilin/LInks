# -*- coding: utf-8 -*-
"""add_published_calibration.py — 在 Discussion 补入已发表数值的校准与不可比声明（R2-M2）"""
from docx import Document

DOC = r'D:\GED_mutation\manuscript_routeA.docx'
DEL = '\u0394'
MINUS = '\u2212'
RHO = '\u03c1'

OLD = ('Published predictors provide useful context but do not isolate this factor because their '
       'architectures, data, and evaluation subsets are different. We did not reproduce ThermoMPNN, '
       'Stability Oracle or any other published structure-based predictor, and the structure-based '
       'arm of this study is therefore the encoder family evaluated here rather than the published '
       'state of the art.')

NEW = ('Published predictors provide useful context but do not isolate this factor because their '
       'architectures, data, and evaluation subsets are different. We did not reproduce ThermoMPNN, '
       'Stability Oracle or any other published structure-based predictor, and the structure-based '
       'arm of this study is therefore the encoder family evaluated here rather than the published '
       'state of the art. For calibration, published structure-based methods report substantially '
       'higher values on their own benchmarks and metrics than the strongest encoder of this ladder '
       'reaches on S669 or ssym; reported examples include a median per-protein Spearman ' + RHO +
       ' of 0.77 for ThermoMPNN on the MegaScale test set. Those quantities are not directly '
       'comparable to ours, because the benchmark, the mutation set, the aggregation (median across '
       'proteins versus pooled across mutations) and the correlation type all differ, and we '
       'therefore present them as context rather than as a ranking. The consequence for this study '
       'is that the null representation effect is established against a model family we trained, '
       'not against the published state of the art, and a representation effect remains possible in '
       'architectures we did not evaluate.')


def main():
    d = Document(DOC)
    n = 0
    for p in d.paragraphs:
        if not p.runs:
            continue
        full = ''.join(r.text for r in p.runs)
        if OLD in full:
            p.runs[0].text = full.replace(OLD, NEW)
            for r in p.runs[1:]:
                r.text = ''
            n += 1
    d.save(DOC)
    print(f'补入已发表数值校准 {n} 处')
    d2 = Document(DOC)
    for p in d2.paragraphs:
        if 'median per-protein Spearman' in p.text:
            i = p.text.find('For calibration')
            print('  ✅', p.text[i:i + 230])


if __name__ == '__main__':
    main()
