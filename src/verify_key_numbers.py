# -*- coding: utf-8 -*-
"""
verify_key_numbers.py — 定向核验：论文关键主张 ↔ 指定源文件
==============================================================
与 number_audit.py 不同，这里为每个关键数字**指定唯一来源**，避免巧合匹配。
"""
import os
import sys

import numpy as np
import pandas as pd

sys.path.insert(0, os.path.dirname(__file__))
from paths import DATA, ROOT

D = str(DATA)
DOCX = os.path.join(str(ROOT), 'manuscript_routeA.docx')


def rd(n):
    p = os.path.join(D, n)
    return pd.read_csv(p) if os.path.exists(p) else None


def pearson(a, b):
    a, b = np.asarray(a, float), np.asarray(b, float)
    return float(np.corrcoef(a, b)[0, 1]) if a.std() > 0 and b.std() > 0 else float('nan')


def check(label, paper_val, actual, tol=0.0015):
    """双向核验：数值须与源文件一致，且须真的出现在论文正文中。

    只看 paper_val 与 actual 是否相符是无效核验——paper_val 是脚本里的常量，
    并未证明论文写的就是这个值。因此额外要求该字符串出现在 DOCX 中。
    """
    try:
        ok = abs(float(paper_val) - float(actual)) <= tol
    except (TypeError, ValueError):
        ok = str(paper_val) == str(actual)
    in_text = str(paper_val) in MANUSCRIPT_TEXT
    mark = '✅' if (ok and in_text) else ('📄' if ok and not in_text else '❌')
    note = '' if in_text else '  ← 未在论文中出现'
    print(f'  {mark} {label:<48} 论文 {paper_val:<12} 实测 {actual}{note}')
    return ok and in_text


def load_manuscript_text():
    """读取 DOCX 正文与表格（body order），用于确认数值确实写在论文里。"""
    global MANUSCRIPT_TEXT
    try:
        from docx import Document
        from docx.oxml.ns import qn
        d = Document(DOCX)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    parts.append(c.text)
        MANUSCRIPT_TEXT = '\n'.join(parts)
        # 归一化 Unicode 减号/连字符，便于匹配
        MANUSCRIPT_TEXT = (MANUSCRIPT_TEXT.replace('\u2212', '-')
                           .replace('\u2013', '-').replace('\u2014', '-')
                           .replace('\u00a0', ' '))
        return True
    except Exception as e:
        print(f'  ⚠ 无法读取论文: {e}')
        MANUSCRIPT_TEXT = ''
        return False


MANUSCRIPT_TEXT = ''



def main():
    if not load_manuscript_text():
        print('❌ 无法读取论文，核验中止')
        return 1
    print(f'  📄 论文文本已加载（{len(MANUSCRIPT_TEXT)} 字符）')
    print()

    results = []

    # ---- S669 / ssym 样本量 ----
    b = rd('benchmarks_s669_clean.csv')
    if b is not None:
        results.append(check('S669 原始总数 543', '543', len(b)))
        if '_node_idx_ok' in b.columns:
            results.append(check('索引对齐通过 512', '512',
                                 int(b['_node_idx_ok'].sum())))
    e = rd('edits_corrected.csv')
    if e is not None:
        # 只取参考阈值 8 Å（文件含 6/7/8/9/10 Å 共 10100 行）
        e8 = e[e['threshold'] == 8.0] if 'threshold' in e.columns else e
        cen = e8[e8['atom_def'] == 'centroid'] if 'atom_def' in e8.columns else e8
        results.append(check('质量受控对 505', '505', len(cen)))
        # 各定义编辑率（8 Å）
        for ad, exp in [('ca', 0.0), ('cb', 9.9), ('centroid', 90.7), ('allatom', 81.6)]:
            sub = e8[e8['atom_def'] == ad]
            if len(sub):
                rate = float((sub['n_edit'] > 0).mean() * 100)
                results.append(check(f'{ad} 编辑率 % (8Å)', exp, round(rate, 1), tol=0.05))

    # ---- 局部性 ----
    loc = rd('locality_corrected.csv')
    if loc is not None and 'dist_broken' in loc.columns:
        results.append(check('断边距离 4.03 Å', '4.03',
                             round(float(loc['dist_broken'].mean()), 2), tol=0.02))
        col = 'dist_kept' if 'dist_kept' in loc.columns else None
        if col:
            results.append(check('保持边距离 15.70 Å', '15.70',
                                 round(float(loc[col].mean()), 2), tol=0.02))

    # ---- Ridge ----
    r = rd('ridge_fixed_results.csv')
    if r is not None:
        r = r.dropna(subset=['r'])
        for _, row in r.iterrows():
            results.append(check(f'Ridge {row["benchmark"]} r',
                                 f'{row["r"]:.3f}', f'{row["r"]:.3f}'))

    # ---- 能力阶梯（S669，审计后的三种子集）----
    lad = rd('ladder_seed_summary_audited.csv')
    if lad is not None:
        for m, lo, hi in [('gnn_global', 0.052, 0.150), ('gnn_local', 0.314, 0.378),
                          ('gnn_edge', 0.340, 0.373), ('deep_gine', 0.314, 0.362),
                          ('egnn', 0.091, 0.131)]:
            s = lad[lad['model'] == m]['seed_mean_r']
            if len(s):
                results.append(check(f'{m} S669 三种子范围下界', f'{lo:.3f}', f'{s.min():.3f}'))
                results.append(check(f'{m} S669 范围上界', f'{hi:.3f}', f'{s.max():.3f}'))

    # ---- ssym 阶梯 ----
    ls = rd('ladder_ssym_results.csv')
    if ls is not None:
        for m, exp in [('gnn_global', 0.147), ('gnn_local', 0.409), ('gnn_edge', 0.452),
                       ('deep_gine', 0.478), ('egnn', 0.235)]:
            s = ls[ls['model'] == m]['r']
            if len(s):
                results.append(check(f'{m} ssym r', f'{exp:.3f}', f'{s.mean():.3f}'))

    # ---- 序列基线 / 融合 ----
    fu = rd('esm2_fusion_results.csv')
    if fu is not None:
        for bm, mdl, exp in [('s669', 'ESM-only', 0.392), ('ssym', 'ESM-only', 0.516),
                             ('s669', 'Graph-only', 0.368), ('ssym', 'Graph-only', 0.460),
                             ('s669', 'ESM+Graph', 0.379), ('ssym', 'ESM+Graph', 0.540)]:
            s = fu[(fu['benchmark'] == bm) & (fu['model'] == mdl)]['r']
            if len(s):
                results.append(check(f'{bm} {mdl} r', f'{exp:.3f}', f'{s.iloc[0]:.3f}'))
        inc = fu[fu['model'].astype(str).str.startswith('increment')]
        for _, row in inc.iterrows():
            results.append(check(f'{row["benchmark"]} 结构增量 Δr',
                                 f'{row["r"]:+.3f}', f'{row["r"]:+.3f}'))
            results.append(check(f'{row["benchmark"]} 增量 CI 下界',
                                 f'{row["ci_low"]:+.3f}', f'{row["ci_low"]:+.3f}'))

    # ---- ESM-2 zero-shot ----
    z = rd('esm2_zeroshot_s669_esm2_650m.csv')
    z2 = rd('esm2_zeroshot_ssym_esm2_650m.csv')
    for df, bm, exp in [(z, 'S669', 0.317), (z2, 'ssym', 0.266)]:
        if df is not None and 'score' in df.columns:
            v = pearson(df['ddg'].values, -df['score'].values)
            results.append(check(f'ESM-2 zero-shot {bm} r', f'{exp:.3f}', f'{v:.3f}',
                                 tol=0.004))

    # ---- 有监督 ESM-2 范围 ----
    sup = rd('esm2_sup_esm2_650m_results.csv')
    if sup is not None:
        for bm, lo, hi in [('s669', 0.348, 0.388), ('ssym', 0.458, 0.493)]:
            s = sup[sup['benchmark'] == bm]['r']
            if len(s):
                results.append(check(f'ESM-2 sup {bm} 下界', f'{lo:.3f}', f'{s.min():.3f}'))
                results.append(check(f'ESM-2 sup {bm} 上界', f'{hi:.3f}', f'{s.max():.3f}'))

    # ---- 接触计数与类型 ----
    dt = rd('directional_type_analysis.csv')
    if dt is not None:
        y = dt['ddg'].values
        for col, exp in [('n_broken', 0.061), ('n_formed', 0.083), ('n_edit', 0.081),
                         ('broken_hydro', 0.135)]:
            if col in dt.columns:
                v = pearson(y, dt[col].values.astype(float))
                results.append(check(f'{col} vs ΔΔG r', f'{exp:.3f}', f'{v:.3f}', tol=0.002))

    # ---- 阈值敏感性 ----
    th = rd('threshold_sensitivity_prediction.csv')
    if th is not None:
        for bm, exp in [('S669', [0.410, 0.386, 0.382, 0.353, 0.353]),
                        ('ssym', [0.381, 0.371, 0.410, 0.359, 0.348])]:
            s = th[th['benchmark'] == bm].sort_values('threshold')['r'].tolist()
            for i, (v, ev) in enumerate(zip(s, exp)):
                results.append(check(f'{bm} {[6,7,8,9,10][i]}Å r', f'{ev:.3f}', f'{v:.3f}'))

    # ---- EGNN 逐种子（架构匹配的三种子）----
    egnn_r = rd('ladder_egnn_legacy_s2024_results.csv')
    lad_raw = rd('ladder_results.csv')
    if egnn_r is not None and lad_raw is not None:
        seeds = lad_raw[(lad_raw['model'] == 'egnn') &
                        (lad_raw['atom_def'] == 'centroid')]['r'].tolist()
        seeds += egnn_r[egnn_r['atom_def'] == 'centroid']['r'].tolist()
        if len(seeds) >= 3:
            results.append(check('EGNN centroid 种子最大值', '0.393', f'{max(seeds):.3f}'))
            results.append(check('EGNN centroid 种子最小值', '-0.081', f'{min(seeds):.3f}'))
            results.append(check('EGNN 三种子均值', '0.131',
                                 f'{np.mean(seeds):.3f}', tol=0.005))
            results.append(check('EGNN 三种子样本SD', '0.241',
                                 f'{np.std(seeds, ddof=1):.3f}', tol=0.005))

    # ---- FoldX 可复现性 ----
    ro = rd('foldx_reproducibility_v2.csv')
    if ro is None:
        ro = rd('foldx_reproducibility.csv')
    if ro is not None:
        results.append(check('FoldX 可复现性行数 3', '3', len(ro)))

    # ---- 独立验证者发现的修正项（防回归）----
    ed = rd('edits_corrected.csv')
    if ed is not None:
        s = ed[(ed.threshold == 8.0) & (ed.atom_def == 'cb')]
        gs = s.gly_switch.astype(bool)
        gly = s[gs]
        ngly = s[~gs]
        results.append(check('甘氨酸 Cβ 有编辑率 72.4%', '72.4',
                             f'{100 * (gly.n_edit > 0).mean():.1f}', tol=0.06))
        results.append(check('非甘氨酸 Cβ 有编辑率 1.8%', '1.8',
                             f'{100 * (ngly.n_edit > 0).mean():.1f}', tol=0.06))

    lo = rd('locality_corrected.csv')
    if lo is not None:
        for pid, mut, exp in [('1R2Y', 'R244E', '19.41'), ('3O39', 'L32P', '17.76'),
                              ('1XZO', 'W36A', '6.78')]:
            r = lo[(lo.pdb_id == pid) & (lo.mut_info.astype(str) == mut)]
            if len(r):
                results.append(check(f'案例 {pid} {mut} 距离', exp,
                                     f'{r.mean_d_broken.iloc[0]:.2f}', tol=0.02))

    sm2 = lad
    if sm2 is not None:
        dgc = sm2[(sm2.model == 'deep_gine') & (sm2.atom_def == 'cb')]
        if len(dgc):
            results.append(check('deep_gine Cβ 三种子均值 0.331', '0.331',
                                 f'{dgc.seed_mean_r.iloc[0]:.3f}', tol=0.002))

    ef2 = rd('ladder_paired_effects_audited.csv')
    if ef2 is not None:
        common = ef2[ef2.comparison.isin(['ca vs centroid', 'cb vs centroid'])]
        for m, exp in [('gnn_global', '0.048'), ('gnn_local', '0.048'),
                       ('gnn_edge', '0.021'), ('deep_gine', '0.043')]:
            g2 = common[common.model == m]
            if len(g2):
                results.append(check(f'{m} 共有对平均|Δr|', exp,
                                     f'{g2.delta_r.abs().mean():.3f}', tol=0.002))
        eg2 = ef2[ef2.model == 'egnn']
        if len(eg2):
            results.append(check('EGNN 平均|Δr| 0.059', '0.059',
                                 f'{eg2.delta_r.abs().mean():.3f}', tol=0.002))

    tr2 = rd('training_merged_noleak_sc.csv')
    if tr2 is not None:
        ms2 = tr2[tr2.source == 'megascale']
        tm2 = tr2[tr2.source == 'thermomutdb']
        n2 = min(len(ms2), len(tm2))
        v2 = pd.concat([ms2.sample(n=n2, random_state=42), tm2])
        results.append(check('训练集 ThermoMutDB 蛋白数 191', '191',
                             str(v2[v2.source == 'thermomutdb'].protein.nunique())))

    # ---- 局域性（排除自身接触后的口径）----
    ns = rd('locality_nonself.csv')
    if ns is not None:
        sub = ns.dropna(subset=['mean_nonself'])
        results.append(check('局域性 非自身断边突变数 234', '234', str(len(sub))))
        results.append(check('局域性 断裂接触平均 7.93 Å', '7.93',
                             f'{sub.mean_nonself.mean():.2f}', tol=0.02))
        results.append(check('局域性 保持接触平均 15.93 Å', '15.93',
                             f'{sub.mean_kept.mean():.2f}', tol=0.02))
        results.append(check('局域性 自身断边数 382', '382',
                             str(int(ns.n_self_edges.sum()))))

    print()
    print('=' * 74)
    n_ok = sum(results)
    print(f'定向核验：{n_ok}/{len(results)} 通过')
    if n_ok < len(results):
        print('❌ 存在不一致，需修正')
    else:
        print('✅ 论文全部关键数字与源文件一致')
    print('=' * 74)
    return 0 if n_ok == len(results) else 1


if __name__ == '__main__':
    sys.exit(main())
