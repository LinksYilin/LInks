# -*- coding: utf-8 -*-
"""fix_figure_block.py — 重建 Figure 3/4/5 的图片与图注配对

现状：
  [159] 失效图片（旧 BLOSUM 图）
  [161] Figure 3 图  [162] Figure 3 图注
  [163] Figure 5 的图（位置错）  [164] Figure 4 图注  [165] Figure 5 图注
目标：
  Figure 3 图 + 图注；Figure 4 图 + 图注；Figure 5 图 + 图注
"""
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

DOC = r'D:\GED_mutation\manuscript_routeA.docx'
FIG = r'D:\GED_mutation\figures\publication'
BLIP = './/{http://schemas.openxmlformats.org/drawingml/2006/main}blip'


def main():
    doc = Document(DOC)
    paras = doc.paragraphs

    def img_paras():
        return [p for p in doc.paragraphs if p._p.findall(BLIP)]

    # 1) 删除失效图片段（位于 Figure 2 图注之后、Figure 3 之前）
    i2 = next(i for i, p in enumerate(doc.paragraphs)
              if p.text.strip().startswith('Figure 2 |'))
    i3 = next(i for i, p in enumerate(doc.paragraphs)
              if p.text.strip().startswith('Figure 3 |'))
    stale = None
    for i in range(i2 + 1, i3):
        if doc.paragraphs[i]._p.findall(BLIP):
            stale = doc.paragraphs[i]
            break
    if stale is not None:
        stale._p.getparent().remove(stale._p)
        print('  已删除失效图片段')
    doc.save(DOC)

    # 2) 把 Figure 4 图注之前的那张图设为 Figure 4；在其后补 Figure 5 的图
    doc = Document(DOC)
    i4 = next(i for i, p in enumerate(doc.paragraphs)
              if p.text.strip().startswith('Figure 4 |'))
    i5 = next(i for i, p in enumerate(doc.paragraphs)
              if p.text.strip().startswith('Figure 5 |'))

    # 图注前一段应是图片段
    def set_image(p, path):
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        p.add_run().add_picture(path, width=Inches(6.4))

    before4 = doc.paragraphs[i4 - 1]
    if before4._p.findall(BLIP):
        set_image(before4, FIG + r'\figure4_information_budget.png')
        print('  Figure 4 图片已就位')
    else:
        newp = OxmlElement('w:p')
        doc.paragraphs[i4]._p.addprevious(newp)
        np_ = Paragraph(newp, doc.paragraphs[i4]._parent)
        np_.add_run().add_picture(FIG + r'\figure4_information_budget.png', width=Inches(6.4))
        print('  Figure 4 图片已插入')

    doc.save(DOC)

    # 3) Figure 5 图注前确保有 Figure 5 的图
    doc = Document(DOC)
    i5 = next(i for i, p in enumerate(doc.paragraphs)
              if p.text.strip().startswith('Figure 5 |'))
    before5 = doc.paragraphs[i5 - 1]
    if not before5._p.findall(BLIP):
        newp = OxmlElement('w:p')
        doc.paragraphs[i5]._p.addprevious(newp)
        np_ = Paragraph(newp, doc.paragraphs[i5]._parent)
        np_.add_run().add_picture(FIG + r'\figure5_contact_diagnostics.png', width=Inches(6.4))
        print('  Figure 5 图片已插入')
    else:
        set_image(before5, FIG + r'\figure5_contact_diagnostics.png')
        print('  Figure 5 图片已更新')
    doc.save(DOC)

    # 4) 报告
    doc = Document(DOC)
    print()
    for i, p in enumerate(doc.paragraphs):
        if p._p.findall(BLIP):
            nxt = doc.paragraphs[i + 1].text.strip()[:50] if i + 1 < len(doc.paragraphs) else ''
            print(f'  [{i}] 图片 -> 下一段: {nxt}')
    print(f'  内嵌图总数: {len(doc.inline_shapes)}')


if __name__ == '__main__':
    main()
