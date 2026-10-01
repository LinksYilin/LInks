# -*- coding: utf-8 -*-
"""audit_docx_health.py — 文档级健康审计

查找同类缺陷：非法属性值、缺失关系、样式问题、空段落堆积、表格异常。
"""
import hashlib
import re
import zipfile
from collections import Counter

from docx import Document
from docx.oxml.ns import qn

DOC = r'D:\GED_mutation\manuscript_routeA.docx'


def main():
    issues = []

    with zipfile.ZipFile(DOC) as z:
        names = z.namelist()
        docxml = z.read('word/document.xml').decode('utf-8')
        styles = z.read('word/styles.xml').decode('utf-8') if 'word/styles.xml' in names else ''
        rels = z.read('word/_rels/document.xml.rels').decode('utf-8')

    # 1) 非法属性值（字面 null 等）
    for label, blob in [('document.xml', docxml), ('styles.xml', styles)]:
        for bad in ['"null"', "'null'", 'w:val="None"', 'w:val="undefined"']:
            c = blob.count(bad)
            if c:
                issues.append(f'{label} 含非法属性值 {bad} × {c}')
        # 空字体名
        for m in re.finditer(r'w:(?:ascii|hAnsi|eastAsia|cs)="([^"]*)"', blob):
            if not m.group(1).strip():
                issues.append(f'{label} 含空字体名')
                break

    # 2) 未被引用的媒体文件
    used = set(re.findall(r'r:embed="(rId\d+)"', docxml))
    rid2target = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', rels))
    media = {n for n in names if n.startswith('word/media/')}
    used_media = {rid2target[r].replace('media/', 'word/media/') for r in used if r in rid2target}
    orphan = media - used_media
    if orphan:
        issues.append(f'未被引用的媒体文件 × {len(orphan)}: {sorted(os.path.basename(o) for o in orphan)}')

    # 3) 引用但缺失的关系
    for r in used:
        if r not in rid2target:
            issues.append(f'关系缺失: {r}')

    # 4) 文档结构
    doc = Document(DOC)
    paras = doc.paragraphs
    styles_used = Counter(p.style.name for p in paras)
    empty = sum(1 for p in paras if not p.text.strip())
    if empty > len(paras) * 0.35:
        issues.append(f'空段落过多: {empty}/{len(paras)}')

    # 5) 表格异常
    for ti, t in enumerate(doc.tables):
        widths = {len(r.cells) for r in t.rows}
        if len(widths) > 1:
            issues.append(f'表 {ti} 行宽不一致: {widths}')

    # 6) 标题层级连续性
    heads = [(i, p.style.name, p.text.strip()) for i, p in enumerate(paras)
             if p.style.name.startswith('Heading')]
    nums = []
    for _, _, t in heads:
        m = re.match(r'^(\d+)(?:\.(\d+))?', t)
        if m:
            nums.append((int(m.group(1)), int(m.group(2) or 0)))
    for a, b in zip(nums, nums[1:]):
        if b[0] == a[0] and b[1] not in (a[1] + 1,):
            issues.append(f'编号跳号: {a} -> {b}')
        if b[0] not in (a[0], a[0] + 1):
            issues.append(f'章节跳号: {a} -> {b}')

    print('=' * 74)
    print('文档健康审计')
    print('=' * 74)
    print(f'段落 {len(paras)}（空 {empty}）| 表 {len(doc.tables)} | 媒体 {len(media)}（未引用 {len(orphan)}）')
    print(f'样式使用: {dict(styles_used.most_common(8))}')
    print()
    if issues:
        print('发现问题:')
        for x in issues:
            print(f'  ⚠ {x}')
    else:
        print('✅ 未发现问题')
    print()
    return len(issues)


if __name__ == '__main__':
    import os
    raise SystemExit(main())
