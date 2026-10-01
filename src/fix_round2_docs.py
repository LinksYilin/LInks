# -*- coding: utf-8 -*-
"""fix_round2_docs.py — 标题页与补充材料的修正"""
import os
import re

from docx import Document

ROOT = r'D:\GED_mutation'
TP = os.path.join(ROOT, '标题页与投稿清单_TitlePage_Checklist.md')
SUP = os.path.join(ROOT, '补充材料_Supplementary.md')

# 从 DOCX 取实际值
d = Document(os.path.join(ROOT, 'manuscript_routeA.docx'))
ps = [p.text.strip() for p in d.paragraphs]
at = ' '.join(t for t in ps if re.match(r'^(Motivation\.|Results\.|Conclusion\.|Availability\.)', t))
NA = len(at.split())
NPARA = len(d.paragraphs)
NFIG = len(d.inline_shapes)
NTAB = len(d.tables)
print(f'实际: 摘要 {NA} 词 | 段落 {NPARA} | 图 {NFIG} | 表 {NTAB}')

TP_RULES = [
    ('Structured abstract, 343 words (journal limit 350).',
     f'Structured abstract, {NA} words (journal limit 350).'),
    ('9 sections: S1 audit trail, S2 provenance, S3 encoder definitions, S4 increment bounds, '
     'S5 ladder and multiplicity, S6 threshold sensitivity, S7 EGNN instability, S8 cross-engine, '
     'S9 reproducibility',
     '10 sections: S1 audit trail, S2 provenance, S3 encoder definitions, S4 increment bounds, '
     'S5 ladder and multiplicity, S6 threshold sensitivity, S7 EGNN instability, S8 cross-engine, '
     'S9 reproducibility, S10 locality statistic'),
    ('Contact-edit diagnostics and cross-engine consistency',
     'Contact-edit diagnostics for the side-chain-centroid definition'),
    ('(DOCX): 16 pages, 230 paragraphs, 5 figures, 1 table',
     f'(DOCX): {NPARA} paragraphs, {NFIG} figures, {NTAB} table'),
    ('`verify_key_numbers.py` — 59/59 checks pass',
     '`verify_key_numbers.py` — 75/75 checks pass (each value is checked both against its source '
     'file and for presence in the manuscript text)'),
    ('`verify_supplementary.py` — 29/29 supplementary numbers trace to a named results file',
     '`verify_supplementary.py` — 29/29 supplementary numbers trace to a named results file'),
]

SUP_RULES = [
    # D-4: S4 +0.010 → +0.009
    ('[−0.036, +0.010]', '[−0.036, +0.009]'),
    ('| −0.036 | +0.010 |', '| −0.036 | +0.009 |'),
    # D-2: "within rounding" 与其自身表格矛盾
    ('The ridge baseline moves only within rounding, confirming that the defect was confined to '
     'the graph-based pathway.',
     'The S669 ridge baseline is unchanged to two decimals (0.392 to 0.390). The ssym entry moves '
     'from 0.316 to 0.305, but those two values come from different evaluation pipelines rather '
     'than from the rebuild: the hydrogen correction cannot affect a ridge model that never reads '
     'the graphs.'),
    # D-3: 239 vs 232
    ('Half of the training set (all MegaScale proteins, 239 graphs) had been built from candidates '
     'that **included hydrogen atoms**',
     'Half of the training set (the 232 MegaScale proteins that survive sampling, 239 graph files '
     'on disk) had been built from candidates that **included hydrogen atoms**'),
    # D-5: S9 文件清单补上架构匹配的 EGNN 文件
    ('Ladder per-seed results and audited aggregates | `data/ladder_results.csv`, '
     '`data/ladder_s2024_results.csv`, `data/ladder_seed_summary_audited.csv`',
     'Ladder per-seed results and audited aggregates | `data/ladder_results.csv` (seeds 42, 123), '
     '`data/ladder_egnn_legacy_s2024_results.csv` (the architecture-matched 508,934-parameter '
     'EGNN seed 2024 run used in the audited aggregates), `data/ladder_seed_summary_audited.csv`. '
     '`data/ladder_s2024_results.csv` holds the rescaled 508,938-parameter variant and is excluded '
     'before aggregation.'),
]


def patch(path, rules, label):
    if not os.path.exists(path):
        print(f'  缺失: {path}')
        return 0
    t = open(path, encoding='utf-8').read()
    n = 0
    for a, b in rules:
        if a in t:
            t = t.replace(a, b)
            n += 1
    open(path, 'w', encoding='utf-8').write(t)
    print(f'  {label}: {n} 处')
    return n


print()
patch(TP, TP_RULES, '标题页')
patch(SUP, SUP_RULES, '补充材料')

t = open(TP, encoding='utf-8').read()
s = open(SUP, encoding='utf-8').read()
print()
print(f'标题页摘要词数正确: {str(NA) in t}')
print(f'标题页 59/59 残留: {"59/59" in t}')
print(f'补充材料 +0.010 残留: {s.count("+0.010")}')
print(f'补充材料 239 说明: {"239 graph files" in s}')
