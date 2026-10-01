# -*- coding: utf-8 -*-
"""fix_round2_manuscript_b.py — 剩余论文缺陷"""
from docx import Document

DOC = r'D:\GED_mutation\manuscript_routeA.docx'
A = '\u03b1'
B = '\u03b2'
DEL = '\u0394'
MIN = '\u2212'

RULES = [
    # A-3: Table 1 单元格 +0.010 → +0.009
    (MIN + '0.013 [' + MIN + '0.036, +0.010]', MIN + '0.013 [' + MIN + '0.036, +0.009]'),
    ('[' + MIN + '0.036, +0.010]', '[' + MIN + '0.036, +0.009]'),

    # E-6/A-1: §4.2 "在每个档位重复四定义比较" → 如实
    ('and repeated the four-definition comparison at each rung: a mean-pooled GCN (5.8 k '
     'parameters),',
     'and repeated the comparison at each rung, with all four graph definitions for the three '
     'lower-capacity encoders and three for the two largest: a mean-pooled GCN (5.8 k parameters),'),

    # A-2: Figure 5 图注 "两个 counts 名义显著" → 与图内标注一致
    ('Two counts are nominally significant, of broken (r = 0.135, 95% CI 0.044 to 0.201, '
     'highlighted) and formed (r = 0.096) hydrophobic contacts; only the broken count survives a '
     'Bonferroni correction across the six tests (threshold 0.0083).',
     'Broken hydrophobic contacts (r = 0.135, 95% CI 0.044 to 0.201, highlighted) are the only '
     'count whose 95% protein-cluster bootstrap interval excludes zero; formed hydrophobic '
     'contacts reach a similar unadjusted correlation (r = 0.096) but their cluster interval '
     'includes zero. Only the broken count survives a Bonferroni correction across the six tests '
     '(threshold 0.0083).'),

    # D-1: 538-33+2=507，不是 505
    ('removed 33 of those while readmitting two, leaving 505 quality-controlled mutation pairs.',
     'removed 33 of those while readmitting two, leaving 507 candidates, of which 505 also carry a '
     'complete contact graph. Contact-change statistics therefore use those 505 pairs.'),
]


def main():
    d = Document(DOC)
    n = 0
    hits = []
    for p in d.paragraphs:
        if not p.runs:
            continue
        full = ''.join(r.text for r in p.runs)
        new = full
        for a, b in RULES:
            if a in new:
                new = new.replace(a, b)
                hits.append(a[:50])
        if new != full:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ''
            n += 1
    # 表格单元格
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    if not p.runs:
                        continue
                    full = ''.join(r.text for r in p.runs)
                    new = full
                    for a, b in RULES:
                        if a in new:
                            new = new.replace(a, b)
                            hits.append(a[:50])
                    if new != full:
                        p.runs[0].text = new
                        for r in p.runs[1:]:
                            r.text = ''
                        n += 1
    d.save(DOC)
    print(f'修改 {n} 处')
    for h in hits:
        print(f'  ✓ {h}')

    d2 = Document(DOC)
    txt = '\n'.join(p.text for p in d2.paragraphs) + '\n' + '\n'.join(
        c.text for t in d2.tables for r in t.rows for c in r.cells)
    print()
    for pat, name in [('+0.010', '残留 +0.010'), ('Two counts are nominally', '残留 Two counts'),
                      ('leaving 505 quality', '残留 505 链'), ('+0.009', '新 +0.009')]:
        print(f'  {name}: {txt.count(pat)}')


if __name__ == '__main__':
    main()
