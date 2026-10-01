# -*- coding: utf-8 -*-
"""apply_3seed_figures.py — 更新 Figure 3 图注/图片 与 Results 4.1 的范围数字"""
import os

from docx import Document
from docx.shared import Inches

DOC = r'D:\GED_mutation\manuscript_routeA.docx'
IMG = r'D:\GED_mutation\figures\publication\figure3_capacity_ladder.png'
DEL = '\u0394'
MINUS = '\u2212'
RHO = '\u03c1'
A = '\u00c5'
AL = '\u03b1'
BE = '\u03b2'

CAP3 = (
    'Figure 3 | A capacity ladder: the representation effect does not grow with encoder size. '
    'a, S669 Pearson r for five encoders spanning 5.8 k to 509 k parameters, each trained on the '
    'audited training set under four graph definitions (C' + AL + ', C' + BE + ', side-chain '
    'centroid, all-atom). Points are three-seed ensembles; the equivariant network uses three '
    'architecture-matched seeds and its point estimates are unstable across them (SD 0.197 on the '
    'side-chain-centroid definition). b, Mean absolute representation effect for the two definition '
    'pairs available at every rung (C' + AL + ' and C' + BE + ' against the side-chain centroid), '
    'as a function of encoder capacity. The four stable rungs span 5.8 k to 460 k parameters and '
    'their mean absolute effect is non-monotone; an exact permutation test on the Spearman '
    'correlation with log parameter count gives ' + RHO + ' = ' + MINUS + '0.80 (exact two-sided '
    'P = 0.333, all 24 permutations). Of the thirteen paired comparisons across the five encoders, '
    'one was significant before correction and none survived Holm correction. The 509 k-parameter '
    'equivariant network (open marker) produced a larger apparent effect but collapsed on a subset '
    'of seeds in every definition tested, so its value reflects optimisation instability rather '
    'than a representation effect. Encoder values are tabulated in Table 1.')

# Results 4.1：[123] 的范围数字
OLD_RANGES = ('the 5.8 k-parameter global-pooling GCN reached 0.026-0.190, the 5.9 k-parameter '
              'mutation-site-pooled GCN 0.310-0.381, the 50 k-parameter edge-aware GINE '
              '0.328-0.418, the 460 k-parameter attention-pooled deep GINE 0.300-0.394, and the '
              '509 k-parameter equivariant network -0.081 to 0.393')
NEW_RANGES = ('the 5.8 k-parameter global-pooling GCN reached 0.052-0.150, the 5.9 k-parameter '
              'mutation-site-pooled GCN 0.314-0.378, the 50 k-parameter edge-aware GINE '
              '0.340-0.373, the 460 k-parameter attention-pooled deep GINE 0.314-0.362 (all '
              'three-seed ensembles), and the 509 k-parameter equivariant network 0.143-0.156 '
              '(three architecture-matched seeds whose individual values span -0.081 to 0.393)')


def main():
    doc = Document(DOC)
    n = 0
    for p in doc.paragraphs:
        if not p.runs:
            continue
        full = ''.join(r.text for r in p.runs)
        new = full
        if full.startswith('Figure 3 | A capacity ladder'):
            new = CAP3
        if OLD_RANGES in new:
            new = new.replace(OLD_RANGES, NEW_RANGES)
        if new != full:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ''
            n += 1

    # 替换 Figure 3 图片
    paras = list(doc.paragraphs)
    i3 = next((i for i, p in enumerate(paras) if p.text.strip().startswith('Figure 3 |')), None)
    BLIP = './/{http://schemas.openxmlformats.org/drawingml/2006/main}blip'
    if i3 is not None and os.path.exists(IMG):
        for j in range(i3 - 1, max(i3 - 4, 0), -1):
            if paras[j]._p.findall(BLIP):
                for r in list(paras[j].runs):
                    r._r.getparent().remove(r._r)
                paras[j].add_run().add_picture(IMG, width=Inches(6.4))
                print(f'  Figure 3 图片已更新（段 {j}）')
                break
    doc.save(DOC)
    print(f'文本更新 {n} 处')


if __name__ == '__main__':
    main()
