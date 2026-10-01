# -*- coding: utf-8 -*-
"""renumber_references.py — 按首次出现顺序重排参考文献编号（Vancouver 格式）

BMC Bioinformatics 采用 Vancouver 格式：参考文献按正文首次出现顺序编号。
当前顺序为 12,8,11,5,4,13,1,10,2,7,15,16,6,14,17,9,3 —— 不符合要求。
"""
import re

from docx import Document

DOC = r'D:\GED_mutation\manuscript_routeA.docx'

d = Document(DOC)
paras = [p.text.strip() for p in d.paragraphs]
iref = next(i for i, t in enumerate(paras) if t == 'References')

# 1) 参考文献列表
refs = {}
for t in paras[iref:]:
    m = re.match(r'^\[(\d+)\]\s+(.+)$', t)
    if m:
        refs[int(m.group(1))] = m.group(2)
print(f'文献 {len(refs)} 条')

# 2) 正文首次出现顺序（含多编号同处的展开）
order = []
for i in range(iref):
    if paras[i] == 'References':
        break
    # 展开 [7], [15] 这类连续写法以及 [9], [3]
    for m in re.finditer(r'\[(\d+)\](?:\s*,\s*\[(\d+)\])*', paras[i]):
        pass
    for m in re.finditer(r'\[(\d+)\]', paras[i]):
        v = int(m.group(1))
        if v not in order:
            order.append(v)
# 处理形如 "[9], [3]" 与 "[7], [15]" 都被上面的 [\d+] 抓到，已覆盖
print(f'首次出现顺序: {order}')
assert set(order) == set(refs), f'顺序集合与文献集合不符: {set(refs) - set(order)}'

mapping = {old: new for new, old in enumerate(order, 1)}
print(f'映射: {mapping}')

# 3) 两阶段替换（避免冲突）
PH = {old: f'\x00{new}\x00' for old, new in mapping.items()}


def renum(text):
    t = text
    for old, ph in PH.items():
        t = re.sub(r'\[' + str(old) + r'\]', ph, t)
    for old, new in mapping.items():
        t = t.replace(f'\x00{new}\x00', f'[{new}]')
    return t


n = 0
for p in d.paragraphs:
    if not p.runs:
        continue
    full = ''.join(r.text for r in p.runs)
    if full.strip() == 'References':
        continue
    new = renum(full)
    if new != full:
        p.runs[0].text = new
        for r in p.runs[1:]:
            r.text = ''
        n += 1

# 4) 重排文献列表
i_ref = next(i for i, p in enumerate(d.paragraphs) if p.text.strip() == 'References')
ref_paras = []
for i in range(i_ref + 1, len(d.paragraphs)):
    if re.match(r'^\[\d+\]', d.paragraphs[i].text.strip()):
        ref_paras.append(d.paragraphs[i])
print(f'文献段落 {len(ref_paras)} 个')

for old, p in zip(order, ref_paras):
    new_num = mapping[old]
    txt = f'[{new_num}] {refs[old]}'
    p.runs[0].text = txt
    for r in p.runs[1:]:
        r.text = ''

d.save(DOC)
print(f'\n重编号 {n} 段正文；文献列表已按新编号排列')

# 5) 复核
d2 = Document(DOC)
ps2 = [p.text.strip() for p in d2.paragraphs]
iref2 = next(i for i, t in enumerate(ps2) if t == 'References')
order2 = []
for i in range(iref2):
    for m in re.finditer(r'\[(\d+)\]', ps2[i]):
        v = int(m.group(1))
        if v not in order2:
            order2.append(v)
print(f'新首次出现顺序: {order2}')
print(f'是否升序: {"✅ 是" if order2 == sorted(order2) else "❌ 否"}')
print()
print('新文献列表:')
for i in range(iref2, len(ps2)):
    if re.match(r'^\[\d+\]', ps2[i]):
        print(f'  {ps2[i][:100]}')
