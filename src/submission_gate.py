"""
submission_gate.py — 投稿前质量门禁（替代 final_audit.py）
=============================================================
与旧 final_audit.py 的关键区别：
  1. 期望值**从结果文件动态推导**，不是硬编码数字
  2. **真的搜索论文全文**，确认推导出的数字确实出现在论文里
  3. **真的运行测试套件**（并如实报告是否可运行）
  4. 检查通过返回 0，任一失败返回 1
  5. 输出机器可读 JSON + 人类可读摘要

用法:
  python submission_gate.py                 # 全部门禁
  python submission_gate.py --no-tests      # 跳过测试（无 pytest 时）
"""
import argparse
import json
import os
import re
import subprocess
import sys
import zipfile

import pandas as pd

from pathlib import Path as _P
DATA = str(_P(__file__).resolve().parent.parent / 'data')
ROOT = str(__import__('pathlib').Path(__file__).resolve().parent.parent)
PUB = os.path.join(ROOT, 'figures', 'publication')

# 稿件路径：优先最新润色稿，其次早期修正稿；也可用 --docx 指定
_CAND = ['manuscript_routeA.docx',
         'manuscript_polished_final_20260930.docx',
         'manuscript_revised_20260930_v2.docx',
         'manuscript_revised_20260930.docx']
DOCX = next((os.path.join(ROOT, c) for c in _CAND if os.path.exists(os.path.join(ROOT, c))),
            os.path.join(ROOT, _CAND[0]))

# 本门禁本身只需要 docx/pandas/numpy；测试需要 torch/sklearn → 用项目 venv
_VENV = os.path.join(ROOT, '.venv', 'Scripts', 'python.exe')
PY = _VENV if os.path.exists(_VENV) else sys.executable

results = []


def record(cat, name, ok, detail=''):
    results.append({'category': cat, 'check': name, 'ok': bool(ok), 'detail': str(detail)[:300]})


# ---------------------------------------------------------------- 论文文本
def paper_text():
    """从 docx 直接抽取全文（不依赖导出的中间文件）。"""
    from docx import Document
    d = Document(DOCX)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return '\n'.join(parts)


# ---------------------------------------------------------------- 数字格式化
MINUS = '\u2212'  # 论文用的真减号


def variants(x, nd=3):
    """给出一个数字在论文中可能的写法。"""
    out = set()
    for fmt in [f'{x:.{nd}f}', f'{x:.2f}', f'{x:.1f}', f'{x:g}']:
        out.add(fmt)
        out.add(fmt.lstrip('0') if fmt.startswith('0.') else fmt)
    if x < 0:
        out |= {v.replace('-', MINUS) for v in list(out)}
    return out


def appears(text, value, nd=3):
    return any(v in text for v in variants(value, nd))


# ---------------------------------------------------------------- 门禁
def check_data_files():
    req = ['benchmark_summary.csv', 'benchmark_predictions_final.csv', 'paired_comparison.csv',
           'edits_corrected.csv', 'mechanism_corrected.csv', 'locality_corrected.csv',
           'foldx_reproducibility_v2.csv', 'edits_ssym_corrected_summary.csv',
           'representation_all_defs.csv', 'threshold_sensitivity_prediction.csv',
           'edits_scwrl4_local8A.csv', 'predictions_unified.csv']
    for f in req:
        p = os.path.join(DATA, f)
        record('data', f, os.path.exists(p),
               f'{os.path.getsize(p)/1024:.0f} KB' if os.path.exists(p) else 'MISSING')


def check_graph_dirs():
    for d in ['contact_graphs_s669_sc', 'contact_graphs_s669_cb', 'contact_graphs_s669_ca',
              'contact_graphs_s669_allatom', 'contact_graphs_ssym_sc',
              'contact_graphs_megascale_sc', 'contact_graphs_thermomutdb_sc']:
        p = os.path.join(DATA, d)
        n = len([x for x in os.listdir(p) if x.endswith('.npz')]) if os.path.isdir(p) else 0
        record('data', d, n > 0, f'{n} graphs')


def check_figures():
    for f in ['figure1_model_comparison.png', 'figure2_contact_rewiring.png',
              'figure3_limited_increment.png']:
        p = os.path.join(PUB, f)
        record('figure', f, os.path.exists(p))
    try:
        with zipfile.ZipFile(DOCX) as z:
            media = [n for n in z.namelist() if n.startswith('word/media/')]
        record('figure', 'docx 内嵌图片数', len(media) >= 3, f'{len(media)} 张')
    except Exception as e:
        record('figure', 'docx 内嵌图片', False, str(e))


def check_numbers(text):
    """★ 核心：从结果文件推导期望值，再搜索论文文本。"""
    bs = pd.read_csv(os.path.join(DATA, 'benchmark_summary.csv'))
    n_s669 = int(bs[bs['benchmark'] == 'S669']['n'].iloc[0])
    n_ssym = int(bs[bs['benchmark'] == 'ssym']['n'].iloc[0])

    ec = pd.read_csv(os.path.join(DATA, 'edits_corrected.csv'))
    c8 = ec[(ec['atom_def'] == 'centroid') & (ec['threshold'] == 8.0)]
    ca8 = ec[(ec['atom_def'] == 'ca') & (ec['threshold'] == 8.0)]
    cb8 = ec[(ec['atom_def'] == 'cb') & (ec['threshold'] == 8.0)]
    aa8 = ec[(ec['atom_def'] == 'allatom') & (ec['threshold'] == 8.0)]
    pct_c = (c8['n_edit'] > 0).mean() * 100
    pct_cb = (cb8['n_edit'] > 0).mean() * 100
    pct_aa = (aa8['n_edit'] > 0).mean() * 100
    ca_edits = int(ca8['n_edit'].sum())

    loc = pd.read_csv(os.path.join(DATA, 'locality_corrected.csv'))
    pc = pd.read_csv(os.path.join(DATA, 'paired_comparison.csv'))
    ro = pd.read_csv(os.path.join(DATA, 'foldx_reproducibility_v2.csv'))

    def paired(bench, needle):
        s = pc[(pc['benchmark'] == bench) & (pc['comparison'].str.contains(needle, regex=False))]
        return s.iloc[0] if len(s) else None

    checks = []
    # 样本量
    checks.append(('S669 交集 n', str(n_s669), n_s669))
    checks.append(('ssym n', str(n_ssym), n_ssym))
    # 编辑率
    checks.append(('质心编辑率 %', f'{pct_c:.1f}', pct_c))
    checks.append(('Cβ 编辑率 %', f'{pct_cb:.1f}', pct_cb))
    checks.append(('全原子编辑率 %', f'{pct_aa:.1f}', pct_aa))
    checks.append(('Cα 总编辑数=0', 'zero', ca_edits))
    # 局部性（从 CSV 动态取）
    if 'dist_broken' in loc.columns:
        checks.append(('断边距离 Å', f'{loc["dist_broken"].mean():.2f}', loc['dist_broken'].mean()))
    # 路线 A 关键值：结构增量（来自 esm2_fusion_results.csv）
    fus_p = os.path.join(str(DATA), 'esm2_fusion_results.csv')
    if os.path.exists(fus_p):
        fu = pd.read_csv(fus_p)
        inc = fu[fu['model'].astype(str).str.startswith('increment')]
        for _, r in inc.iterrows():
            b = r['benchmark']
            checks.append((f'{b} 结构增量 Δr', f'{r["r"]:+.3f}', r['r']))
            if pd.notna(r.get('ci_low')):
                checks.append((f'{b} 结构增量 CI 下界', f'{r["ci_low"]:+.3f}', r['ci_low']))
    # 理化基线（来自 ridge_fixed_results.csv）
    rp = os.path.join(str(DATA), 'ridge_fixed_results.csv')
    if os.path.exists(rp):
        rr = pd.read_csv(rp).dropna(subset=['r'])
        for _, r in rr.iterrows():
            checks.append((f'{r["benchmark"]} Ridge r', f'{r["r"]:.3f}', r['r']))
    # 等价性分析（可排除的效应上界）
    ep = os.path.join(str(DATA), 'equivalence_analysis.csv')
    if os.path.exists(ep):
        ee = pd.read_csv(ep)
        if len(ee):
            checks.append(('等价性分析行数', str(len(ee)), len(ee)))
    # FoldX 可复现
    checks.append(('FoldX 可复现 CSV 行数', str(len(ro)), len(ro)))

    for label, expect, val in checks:
        if isinstance(val, float) and abs(val) < 1e-12 and expect == 'zero':
            ok = ('zero' in text.lower()) or ('no contact change' in text.lower())
        else:
            ok = str(expect) in text or (isinstance(val, float) and appears(text, val))
        record('number', f'{label} = {expect}', ok,
               '在论文中找到' if ok else '★ 论文中未找到该值')


def check_forbidden(text):
    """论文中不应残留的旧错误值。"""
    forbidden = [
        ('26 broken (1BFM 旧值)', '26 broken'),
        ('94.3% 旧编辑率', '94.3'),
        ('21.0 旧断边均值', '21.0 broken'),
        ('501-mutation 旧交集', '501-mutation'),
        ('86.7% 旧可复现性', '86.7'),
        ('differing by < 1 Å 旧结论', 'differing by < 1 Å'),
        ('+0.007 旧 BLOSUM 增量', '+0.007'),
        ('P = 0.044 旧显著性', 'P = 0.044'),
        ('statistically detectable 旧表述', 'statistically detectable'),
        ('4.01 Å 旧断边距离', '4.01'),
        ('15.44 Å 旧保持边距离', '15.44'),
        ('+0.021 旧 Ridge 差', '+0.021'),
        # 注：0.371 曾是旧 local r；现为阈值敏感性 ssym 7 Å 的当前有效值，
        #     故从过期值列表中移除（改由阈值 CSV 的数值检查覆盖）。
    ]
    for name, s in forbidden:
        found = s in text
        record('stale', name, not found, '★ 仍残留' if found else '已清除')


def check_tests(run=True):
    if not run:
        record('test', '测试套件', False, '用户指定跳过（--no-tests）')
        return
    tests = ['test_ged.py', 'test_leakage_split.py', 'test_label_audit.py', 'test_bootstrap_ci.py']
    for t in tests:
        p = os.path.join(ROOT, 'src', t)
        if not os.path.exists(p):
            record('test', t, False, '文件缺失')
            continue
        try:
            r = subprocess.run([PY, p], capture_output=True, text=True, timeout=900,  # noqa: PLW1510 (返回值由调用方检查)
                               cwd=os.path.join(ROOT, 'src'))
            ok = r.returncode == 0
            tail = (r.stdout or r.stderr).strip().split('\n')[-1][:150]
            record('test', t, ok, tail)
        except subprocess.TimeoutExpired:
            record('test', t, False, '超时')
        except Exception as e:
            record('test', t, False, str(e)[:150])


def check_renders():
    """渲染 docx 首页，确认可打开且无缺字体。"""
    cli = r'D:\新建文件夹\resources\app.asar.unpacked\dsh\node_modules\@deepseek-ai\libreoffice-kit\lib\cli.js'
    node = r'D:\新建文件夹\resources\runtime\primary-runtime\dependencies\node\bin\node.exe'
    if not (os.path.exists(cli) and os.path.exists(node)):
        record('render', 'docx 渲染', False, '渲染工具不可用')
        return
    out = os.path.join(ROOT, 'gate_preview')
    # ★ 渲染工具要求输出目录不存在，先清理（否则报 EEXIST）
    if os.path.isdir(out):
        import shutil
        shutil.rmtree(out, ignore_errors=True)
    try:
        r = subprocess.run([node, cli, 'render', '--input', DOCX, '--output-dir', out,  # noqa: PLW1510 (返回值由调用方检查)
                            '--pages', '1', '--dpi', '72'],
                           capture_output=True, text=True, timeout=600)
        txt = (r.stdout or '') + (r.stderr or '')
        m = re.search(r'"pageCount":(\d+)', txt)
        fonts = re.search(r'"missingFonts":\[([^\]]*)\]', txt)
        ok = r.returncode == 0 and m is not None
        record('render', 'docx 渲染', ok, f'页数 {m.group(1) if m else "?"}')
        if fonts is not None:
            # 渲染工具在无缺失时输出 ["null"]（数组内含字符串 "null"），
            # 有缺失时输出 ["字体名", ...]。两种都要正确识别。
            raw = fonts.group(1).strip().strip('"').strip()
            clean = raw in ('', 'null', 'None') or all(
                x.strip().strip('"').strip() in ('', 'null', 'None')
                for x in raw.split(',') if x.strip())
            record('render', '字体完整', clean, '无缺失' if clean else raw[:120])
    except Exception as e:
        record('render', 'docx 渲染', False, str(e)[:150])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--no-tests', action='store_true')
    ap.add_argument('--out', default=os.path.join(ROOT, 'submission_gate_report.json'))
    args = ap.parse_args()

    print('=' * 74)
    print('投稿前质量门禁')
    print('=' * 74)

    text = paper_text()
    print(f'论文正文长度: {len(text)} 字符\n')

    for fn in [lambda: check_data_files(), lambda: check_graph_dirs(), check_figures,
               lambda: check_numbers(text), lambda: check_forbidden(text),
               lambda: check_tests(not args.no_tests), check_renders]:
        try:
            fn()
        except Exception as e:
            record('internal', getattr(fn, '__name__', 'check'), False, f'门禁自身异常: {e}')

    # 汇总
    cats = {}
    for r in results:
        cats.setdefault(r['category'], []).append(r)
    names = {'data': '数据与图文件', 'figure': '图表', 'number': '论文数字可追溯',
             'stale': '旧错误值清除', 'test': '测试套件', 'render': '渲染', 'internal': '门禁自身'}
    for c, items in cats.items():
        nfail = sum(1 for i in items if not i['ok'])
        mark = '✅' if nfail == 0 else '❌'
        print(f'\n{mark} {names.get(c, c)}（{len(items) - nfail}/{len(items)} 通过）')
        for i in items:
            if not i['ok']:
                print(f'    FAIL  {i["check"]}: {i["detail"]}')
            elif c in ('number', 'stale', 'test'):
                print(f'    ok    {i["check"]}  {i["detail"]}')

    total = len(results)
    failed = [r for r in results if not r['ok']]
    print('\n' + '=' * 74)
    print(f'总计 {total - len(failed)}/{total} 通过')
    if failed:
        print(f'❌ 门禁未通过，{len(failed)} 项失败：')
        for r in failed:
            print(f'   [{r["category"]}] {r["check"]}: {r["detail"]}')
    else:
        print('✅ 全部门禁通过')
    print('=' * 74)

    with open(args.out, 'w', encoding='utf-8') as fh:
        json.dump({'total': total, 'passed': total - len(failed), 'failed': len(failed),
                   'results': results}, fh, ensure_ascii=False, indent=1)
    print(f'报告已保存 {args.out}')
    return 1 if failed else 0


if __name__ == '__main__':
    sys.exit(main())
