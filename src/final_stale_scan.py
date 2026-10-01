# -*- coding: utf-8 -*-
"""final_stale_scan.py — 全面扫描所有已知被取代的数值"""
import re

from docx import Document

DOC = r'D:\GED_mutation\manuscript_routeA.docx'
SUP = r'D:\GED_mutation\补充材料_Supplementary.md'

# (模式, 说明, 是否允许)
STALE = [
    # S10 对照表中故意保留的旧口径（用于展示定义敏感性），允许出现在补充材料
    (r'4\.03 \u00c5', '旧局域性 4.03 Å（S10 对照表）', True),
    (r'15\.70 \u00c5', '旧保持接触 15.70 Å（S10 对照表）', True),
    (r'two seeds per rung', '旧两种子局限', False),
    (r'used two seeds \(42 and 123\)', '旧两种子说明', False),
    (r'SD 0\.197', '旧总体 SD', False),
    (r'74\.6%', '旧甘氨酸率', False),
    (r'6\.8% of others', '旧非甘氨酸率', False),
    (r'0\.362 on the same', '旧 deep_gine Cβ', False),
    (r'mean absolute effect of 0\.061', '旧 EGNN 效应', False),
    (r'0\.066, 0\.048, 0\.019', '旧全集平均', False),
    (r'0\.071 \(5\.8 k\)', '旧重复趋势段', False),
    (r'20,000 permutations', '旧置换次数', False),
    (r'r = 0\.380 for the side-chain', '旧两种子 0.380', False),
    (r'r = 0\.306', '旧 ssym ridge', False),
    (r'6\.5 \u00c5 from the mutated', '旧案例距离', False),
    (r'maximum 0\.300 \u00c5', '旧 SCWRL4 位移', False),
    (r'mean 0\.096 \u00c5', '旧 SCWRL4 均值', False),
    (r'from 2\.1 to 26\.3', '旧氢修复对比', False),
    (r'188 from ThermoMutDB', '旧训练蛋白数', False),
    (r'nested subsets', '旧嵌套声明', False),
    (r'three-seed ensemble', '旧 ensemble 误标', False),
    (r'cross-engine consistency\. a', '旧 Figure 5 标题', False),
    (r'seed-specific uncertainty', '旧 Figure 1 描述', False),
    (r'512 \u2212 33', '旧算式', False),
    (r'33 mutation pairs\)\. Contact-change', '旧 33 归属', False),
    # 应保留
    (r'\+0\.009', '新 S669 上界', True),
    (r'\+0\.070', 'ssym 上界', True),
    (r'sample SD 0\.241', '新样本 SD', True),
    (r'7\.93', '新局域性', True),
    (r'mean of three per-seed', '新 ensemble 标签', True),
]

d = Document(DOC)
doc_txt = '\n'.join(p.text for p in d.paragraphs) + '\n' + '\n'.join(
    c.text for t in d.tables for r in t.rows for c in r.cells)
sup_txt = open(SUP, encoding='utf-8').read()

print('=' * 76)
print('陈旧值全面扫描')
print('=' * 76)
bad = 0
for pat, name, allowed in STALE:
    nd = len(re.findall(pat, doc_txt, re.I))
    ns = len(re.findall(pat, sup_txt, re.I))
    if not allowed and (nd or ns):
        print(f'  ❌ {name:<26} 稿件 {nd}  补充材料 {ns}')
        for m in re.finditer(r'.{0,80}' + pat + r'.{0,80}', doc_txt, re.I):
            print(f'       …{m.group(0)}…')
            break
        bad += nd + ns
    elif allowed and (nd or ns):
        print(f'  ✅ {name:<26} 稿件 {nd}  补充材料 {ns}')

print()
print(f'需要处理: {bad} 处' if bad else '✅ 无陈旧值残留')
