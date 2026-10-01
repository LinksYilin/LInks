# -*- coding: utf-8 -*-
"""fix_r1m3_v2.py — 用锚点切片替换（避免 Unicode 减号/连字符不匹配）"""
from docx import Document

DOC = r'D:\GED_mutation\manuscript_routeA.docx'

START = 'Stratifying by contact type did produce one significant quantity'
END = 'contacts were all non-significant.'
NEW = ('Stratifying by contact type produced two nominally significant quantities among the six '
       'tests: the number of broken hydrophobic contacts (r = 0.135, 95% CI 0.044 to 0.201, '
       'nominal P = 0.0024) and the number of formed hydrophobic contacts (r = 0.096, nominal '
       'P = 0.032). Broken electrostatic (r = 0.047), other broken (r = 0.027), formed '
       'electrostatic (r = -0.046) and other formed (r = 0.063) contacts were not significant. '
       'Only the broken hydrophobic count survives a Bonferroni correction across the six tests '
       '(threshold 0.0083), so we treat it as the one quantity that withstands multiplicity '
       'control rather than as a feature selected from six.')


def main():
    d = Document(DOC)
    n = 0
    for p in d.paragraphs:
        if not p.runs:
            continue
        full = ''.join(r.text for r in p.runs)
        i = full.find(START)
        if i == -1:
            continue
        j = full.find(END, i)
        if j == -1:
            print('  ⚠ 找到起点但未找到终点')
            continue
        j += len(END)
        new = full[:i] + NEW + full[j:]
        p.runs[0].text = new
        for r in p.runs[1:]:
            r.text = ''
        n += 1
    d.save(DOC)
    print(f'修正 {n} 处')
    d2 = Document(DOC)
    for i, p in enumerate(d2.paragraphs):
        if 'the only contact-derived feature to exclude zero' in p.text:
            print(f'  ❌ 残留于 [{i}]')
        if 'two nominally significant quantities among the six tests' in p.text:
            print(f'  ✅ 已修 [{i}]')


if __name__ == '__main__':
    main()
