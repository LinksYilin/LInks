# -*- coding: utf-8 -*-
"""refresh_embedded_figures.py — 用磁盘上的最新 PNG 替换 DOCX 中嵌入的图

背景：修改绘图脚本后重新生成了 PNG，但 DOCX 里嵌入的是旧副本。
本脚本按 sha256 比对，仅替换内容确实变化的图。
"""
import hashlib
import os
import re

from docx import Document
from docx.oxml.ns import qn

DOCX = r'D:\GED_mutation\manuscript_routeA.docx'
FIG = r'D:\GED_mutation\figures\publication'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'

# 论文图：段落标题 -> 磁盘文件
WANT = [
    ('Figure 1 |', 'figure1_study_design.png'),
    ('Figure 2 |', 'figure2_contact_rewiring.png'),
    ('Figure 3 |', 'figure3_capacity_ladder.png'),
    ('Figure 4 |', 'figure4_information_budget.png'),
    ('Figure 5 |', 'figure5_contact_diagnostics.png'),
]

d = Document(DOCX)
# 建立 sha -> 文件名
disk = {}
for f in os.listdir(FIG):
    if f.endswith('.png') and f.startswith('figure'):
        disk[hashlib.sha256(open(os.path.join(FIG, f), 'rb').read()).hexdigest()] = f

changed = 0
for i, p in enumerate(d.paragraphs):
    cap = p.text.strip()
    match = next((f for pre, f in WANT if cap.startswith(pre)), None)
    if not match:
        continue
    # 图注前一段应是图片段
    for j in range(i - 1, max(i - 4, -1), -1):
        blips = d.paragraphs[j]._p.findall('.//' + qn('a:blip'))
        if not blips:
            continue
        rid = blips[0].get(R + 'embed')
        try:
            cur = hashlib.sha256(d.part.rels[rid].target_part.blob).hexdigest()
        except Exception:
            continue
        want_sha = hashlib.sha256(open(os.path.join(FIG, match), 'rb').read()).hexdigest()
        if cur == want_sha:
            print(f'  {match}: 已是最新')
            break
        # 替换
        from docx.shared import Inches
        para = d.paragraphs[j]
        for r in list(para.runs):
            r._r.getparent().remove(r._r)
        para.add_run().add_picture(os.path.join(FIG, match), width=Inches(6.4))
        print(f'  {match}: 已更新（段 {j}）')
        changed += 1
        break

d.save(DOCX)
print(f'\n共更新 {changed} 张图')
