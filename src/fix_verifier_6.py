# -*- coding: utf-8 -*-
"""fix_verifier_6.py — 统一 S669 结构增量的上界为 +0.009（存储值 0.009460）"""
import re

from docx import Document

DOC = r'D:\GED_mutation\manuscript_routeA.docx'
DEL = '\u0394'

# 只替换与 S669 增量上界相关的 +0.010 / 0.01
PATTERNS = [
    (r'above \+0\.010 on S669', 'above +0.009 on S669'),
    (r'excluded below \u22120\.036 and above \+0\.010', 'excluded below \u22120.036 and above +0.009'),
    (r'95% CI \u22120\.036 to \+0\.010', '95% CI \u22120.036 to +0.009'),
    (r'\u22120\.036 to \+0\.010', '\u22120.036 to +0.009'),
    (r'no increment larger than about \+0\.01 on S669', 'no increment larger than about +0.009 on S669'),
    (r'above about \+0\.01 on S669', 'above about +0.009 on S669'),
]


def main():
    d = Document(DOC)
    n = 0
    for p in d.paragraphs:
        if not p.runs:
            continue
        full = ''.join(r.text for r in p.runs)
        new = full
        for pat, rep in PATTERNS:
            new2 = re.sub(pat, rep, new)
            if new2 != new:
                n += 1
                new = new2
        if new != full:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ''
    d.save(DOC)
    print(f'替换 {n} 处')

    d2 = Document(DOC)
    txt = '\n'.join(p.text for p in d2.paragraphs)
    print()
    for pat, name in [(r'\+0\.010', '残留 +0.010'),
                      (r'\+0\.009', '新 +0.009'),
                      (r'\+0\.070', 'ssym +0.070（应保留）')]:
        print(f'  {name}: {len(re.findall(pat, txt))}')


if __name__ == '__main__':
    main()
