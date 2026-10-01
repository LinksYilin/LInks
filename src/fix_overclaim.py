# -*- coding: utf-8 -*-
"""
fix_overclaim.py — R3-m13：把"no incremental information"改为有界表述
=====================================================================
审稿人 3 指出：CI 上界 +0.037 与论文它处报告的增量同量级，
"added no incremental information" 强于区间所能支持。
"""
from docx import Document

DOC = r'D:\GED_mutation\manuscript_routeA.docx'
D = '\u0394'
MINUS = '\u2212'

A1 = ('added no incremental information over them (' + D + 'r = ' + MINUS +
      '0.009, 95% CI ' + MINUS + '0.047 to +0.037)')
B1 = ('added no increment above about +0.04 over them (' + D + 'r = ' + MINUS +
      '0.009, 95% CI ' + MINUS + '0.047 to +0.037)')

A2 = ('added no incremental information over those features (' + D + 'r = ' + MINUS +
      '0.009, 95% CI ' + MINUS + '0.047 to +0.037)')
B2 = ('added no increment above about +0.04 over those features (' + D + 'r = ' + MINUS +
      '0.009, 95% CI ' + MINUS + '0.047 to +0.037)')

FIX = [(A1, B1), (A2, B2)]


def main():
    d = Document(DOC)
    n = 0
    for p in d.paragraphs:
        if not p.runs:
            continue
        full = ''.join(r.text for r in p.runs)
        new = full
        for a, b in FIX:
            if a in new:
                new = new.replace(a, b)
                n += 1
        if new != full:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ''
    d.save(DOC)
    print(f'修正 {n} 处')
    d2 = Document(DOC)
    for p in d2.paragraphs:
        if 'no increment above about +0.04' in p.text:
            i = p.text.find('no increment above about +0.04')
            print('  ✅', p.text[max(0, i - 90):i + 90])


if __name__ == '__main__':
    main()
