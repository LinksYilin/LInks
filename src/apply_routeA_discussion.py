# -*- coding: utf-8 -*-
"""
apply_routeA_discussion.py — Discussion 重构（保留有效内容，更新失效数字，插入新发现）
=====================================================================================
失效数字来源：旧 BLOSUM/消融数字基于含氢 bug 的训练数据。
新增内容：能力阶梯、序列基线增量、信息预算、接触类型分解。
"""
import os

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

DOC = r'D:\GED_mutation\manuscript_routeA.docx'


def insert_after(par, text, style=None):
    new_p = OxmlElement('w:p')
    par._p.addnext(new_p)
    np = Paragraph(new_p, par._parent)
    if style is not None:
        try:
            np.style = style
        except Exception:
            pass
    np.add_run(text)
    return np


def set_text(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.add_run(text)


NEW_52_HEAD = '5.2 Why representational sensitivity does not translate into predictive gain'
NEW_52_BODY = (
    'Three observations explain the decoupling. First, the additional contact information a '
    'sensitive definition supplies is largely redundant with what the model already knows: the '
    'mutation itself fixes the physicochemical change, and five such features reach r = 0.390 on '
    'S669, comparable to every graph encoder tested. The one structural quantity that correlates '
    'with \u0394\u0394G on its own, the count of broken hydrophobic contacts (r = 0.135, 95% CI '
    '0.044 to 0.201), adds no incremental information over those features (cross-validated \u0394r '
    '= -0.009, 95% CI -0.047 to +0.037); it is redundant with the hydrophobicity change the '
    'physicochemical block already contains. Second, the remaining contact-derived quantities are '
    'weak and mutually correlated: net change and broken-to-formed ratio were non-significant '
    '(r = -0.013 and 0.031), and breaking and forming are themselves correlated at r = 0.42, so '
    'directional summaries cancel rather than sharpen the signal. Third, supplying more informative '
    'edges does not by itself change what an aggregation function extracts: encoders with explicit '
    'edge features did not outperform those without them at matched capacity.')

NEW_53_BODY = (
    'A single architecture cannot separate a statement about representations from a statement '
    'about that architecture, so we repeated the four-definition comparison at five encoder '
    'capacities spanning 5.8 k to 509 k parameters. Across the four stable encoders, none of the '
    'eight paired definition comparisons reached significance, and the mean absolute effect did not '
    'grow with capacity (0.052, 0.052, 0.031 and 0.039). This is the pattern that would be absent '
    'if limited encoder strength were the reason the representation effect was not detected. The '
    '509 k-parameter equivariant network produced a larger apparent effect, but it collapsed on a '
    'subset of seeds in every definition (r = 0.393 and -0.081 for two seeds on the '
    'side-chain-centroid graphs), so that comparison reflects optimisation instability rather than '
    'a representation effect and is reported as such. '
    'The same control applies to the sequence comparison. The graph branch used for the fusion '
    'reached r = 0.368 on S669 and 0.460 on ssym on its own, close to the ESM-2 baseline '
    '(0.392 and 0.516), so the absence of a structure increment (S669 \u0394r = -0.013, 95% CI '
    '-0.036 to +0.010; ssym \u0394r = +0.025, 95% CI -0.007 to +0.070) is not attributable to an '
    'under-powered graph branch.')

NEW_179 = (
    'Predictive ablation sharpens this interpretation further. Using the audited training set and '
    'the side-chain-centroid definition, a five-feature physicochemical ridge regression reached '
    'r = 0.390 on S669 and 0.306 on ssym, while the contact-count features contributed a '
    'cross-validated increment of \u0394r = -0.009 (95% CI -0.047 to +0.037) once the '
    'physicochemical block was present. Adding five graph-derived structural descriptors to the '
    'physicochemical features raised r from 0.390 to 0.427 on S669 and from 0.306 to 0.408 on ssym, '
    'an increment that is small in absolute terms and does not survive the addition of a sequence '
    'baseline. The typed contact counts behave the same way, which is why we report the '
    'univariate significance and the multivariate redundancy together rather than either alone.')

NEW_LIM_HEAD = '5.5 Limitations'
NEW_LIM_BODY = (
    'Five boundaries define the scope of these conclusions. First, the predictive results span '
    'encoders from 5.8 k to 509 k parameters; a representation effect in larger or differently '
    'biased architectures is not excluded. Second, ESM-2 was pre-trained on UniRef, so the '
    'sequence baseline may have seen proteins related to the test sets; the zero-shot variant '
    'carries no \u0394\u0394G supervision and is less affected, and it scores below the graph '
    'encoders on the same mutations. Third, S669 guided model development and is exploratory, '
    'while ssym is development-independent but contains only 15 proteins, which widens its '
    'intervals; the structure increment there can only be bounded above \u2248 0.07. Fourth, '
    'absolute contact counts depend on the structure-modelling engine, and FoldX and SCWRL4 differ '
    'in whether they rebuild C\u03b2 coordinates, which changes the reported edit rate for the '
    'C\u03b2 definition. Fifth, the equivariant encoder was unstable across seeds, so the highest '
    'capacity rung is represented by the 460 k-parameter attention-pooled GINE network.')


def main():
    doc = Document(DOC)
    paras = doc.paragraphs
    ps = [p.text.strip() for p in paras]

    def idx_of(prefix, start=0):
        for i in range(start, len(ps)):
            if ps[i].startswith(prefix):
                return i
        return None

    # 1) 5.1 标题改名
    i = idx_of('5.1 ')
    if i is not None:
        set_text(paras[i], '5.1 What changes structurally, and what that does not imply')
        print('  [1] 5.1 标题已改')

    # 2) 在 5.1 正文最后一段之后插入新的 5.2
    i52 = idx_of('5.2 ')
    if i52 is not None:
        anchor = paras[i52 - 1]
        h = insert_after(anchor, NEW_52_HEAD, style=paras[i52].style)
        insert_after(h, NEW_52_BODY)
        print('  [2] 新 5.2 已插入')

    # 重新读取（插入后索引变化）
    doc.save(DOC)
    doc = Document(DOC)
    paras = doc.paragraphs
    ps = [p.text.strip() for p in paras]

    def idx2(prefix, start=0):
        for i in range(start, len(ps)):
            if ps[i].startswith(prefix):
                return i
        return None

    # 3) 旧 5.2 → 5.3（保留其内容）
    i = idx2('5.2 Contact changes are reproducible')
    if i is not None:
        set_text(paras[i], '5.3 Contact changes are reproducible but model-dependent structural '
                           'descriptions')
        print('  [3] 旧 5.2 → 5.3')

    # 4) 更新失效的消融段（以 "Predictive ablation sharpens" 开头）
    for p in paras:
        if p.text.strip().startswith('Predictive ablation sharpens'):
            set_text(p, NEW_179)
            print('  [4] 消融段数字已更新')
            break

    # 5) 旧 5.3 → 5.4，并插入能力阶梯段
    i = idx2('5.3 Implications for representation-layer work')
    if i is not None:
        set_text(paras[i], '5.4 Where the signal is, and what to do about it')
        # 在其后插入能力阶梯内容（作为 5.4 的第一段）
        insert_after(paras[i], NEW_53_BODY)
        print('  [5] 旧 5.3 → 5.4，并插入能力阶梯段')

    # 6) 旧 5.4 → 5.5，并更新
    doc.save(DOC)
    doc = Document(DOC)
    paras = doc.paragraphs
    for p in paras:
        t = p.text.strip()
        if t.startswith('5.4 Limitations'):
            set_text(p, NEW_LIM_HEAD)
            print('  [6] 旧 5.4 → 5.5')
        elif t.startswith('The conclusions have five main boundaries'):
            set_text(p, NEW_LIM_BODY)
            print('  [6] 限制段已更新')

    doc.save(DOC)
    print(f'\n已保存 {DOC}')


if __name__ == '__main__':
    main()
