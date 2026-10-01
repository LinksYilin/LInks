# -*- coding: utf-8 -*-
"""fix_null_font.py — 修复标题字体被设为字面字符串 'null' 的缺陷

25 个 Heading run 的 rFonts 全部为 'null'（ascii/eastAsia/hAnsi/cs），
导致渲染器回退到符号字体，标题显示成希腊字母样式。
改为与正文一致的字体。
"""
import re

from docx import Document
from docx.oxml.ns import qn

DOC = r'D:\GED_mutation\manuscript_routeA.docx'
FONT = 'Times New Roman'


def main():
    d = Document(DOC)

    # 先探测正文字体，保持全文一致
    body_fonts = {}
    for p in d.paragraphs:
        if p.style.name.startswith('Heading'):
            continue
        for r in p.runs:
            if r.font.name and r.font.name != 'null':
                body_fonts[r.font.name] = body_fonts.get(r.font.name, 0) + 1
    print('正文字体分布:', dict(sorted(body_fonts.items(), key=lambda x: -x[1])[:4]))
    target = max(body_fonts, key=body_fonts.get) if body_fonts else FONT
    print(f'采用: {target}')

    n = 0
    for p in d.paragraphs:
        for r in p.runs:
            rPr = r._r.find(qn('w:rPr'))
            if rPr is None:
                continue
            rf = rPr.find(qn('w:rFonts'))
            if rf is None:
                continue
            changed = False
            for key in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
                k = qn(key)
                if rf.get(k) == 'null':
                    rf.set(k, target)
                    changed = True
            if changed:
                n += 1
    d.save(DOC)
    print(f'修复 {n} 个 run 的字体')

    # 验证
    d2 = Document(DOC)
    left = sum(1 for p in d2.paragraphs for r in p.runs if r.font.name == 'null')
    print(f'剩余 font.name==null: {left}')


if __name__ == '__main__':
    main()
