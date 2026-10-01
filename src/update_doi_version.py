# -*- coding: utf-8 -*-
"""update_doi_version.py — 把版本号更新为 v1.0.2 并核验概念 DOI"""
import os
import re

from docx import Document

ROOT = r'D:\GED_mutation'
DOC = os.path.join(ROOT, 'manuscript_routeA.docx')
CONCEPT = '10.5281/zenodo.23086956'
V_OLD = '10.5281/zenodo.23086957'      # v1.0.1（过时）
V_NEW = '10.5281/zenodo.23087054'      # v1.0.2（当前）
URL_OLD = 'https://zenodo.org/records/23086957'
URL_NEW = 'https://zenodo.org/records/23087054'

# 论文：把 records 链接指向最新版
d = Document(DOC)
n = 0
for p in d.paragraphs:
    if not p.runs:
        continue
    full = ''.join(r.text for r in p.runs)
    new = full.replace(URL_OLD, URL_NEW)
    if new != full:
        p.runs[0].text = new
        for r in p.runs[1:]:
            r.text = ''
        n += 1
d.save(DOC)
print(f'论文链接更新: {n} 处')

# markdown 材料
files = ['补充材料_Supplementary.md', '标题页与投稿清单_TitlePage_Checklist.md',
         'CITATION.cff', '如何拿到DOI_三步指南.md', '发布与归档操作指南.md']
for f in files:
    p = os.path.join(ROOT, f)
    if not os.path.exists(p):
        continue
    t = open(p, encoding='utf-8').read()
    o = t
    t = t.replace(URL_OLD, URL_NEW).replace(V_OLD, V_NEW)
    t = t.replace('version v1.0.1', 'version v1.0.2')
    t = t.replace('v1.0.1', 'v1.0.2')
    if t != o:
        open(p, 'w', encoding='utf-8').write(t)
        print(f'  {f}: 已更新')

print()
print('=== 校验 ===')
d2 = Document(DOC)
txt = '\n'.join(p.text for p in d2.paragraphs)
print(f'论文概念 DOI ({CONCEPT}): {txt.count(CONCEPT)} 处')
print(f'论文旧版本 DOI 残留: {txt.count(V_OLD)}')
print(f'论文新版链接 ({URL_NEW}): {txt.count(URL_NEW)}')
for f in ['补充材料_Supplementary.md', '标题页与投稿清单_TitlePage_Checklist.md', 'CITATION.cff']:
    t = open(os.path.join(ROOT, f), encoding='utf-8').read()
    print(f'  {f}: 概念 {t.count(CONCEPT)} | 旧版残留 {t.count(V_OLD)}')
