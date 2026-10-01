# -*- coding: utf-8 -*-
"""fix_sd_convention.py — 统一 EGNN 种子 SD 为样本 SD（ddof=1）

ddof=0（总体 SD）：原版 0.1968，改进版 0.2155
ddof=1（样本 SD）：原版 0.2410，改进版 0.2639
论文原用 ddof=0 的 0.197/0.216，与「SD across three seeds」的常规报告口径不符。
"""
import numpy as np
from docx import Document

DOC = r'D:\GED_mutation\manuscript_routeA.docx'

orig = [0.3932, -0.0805, 0.0794]
impr = [-0.1815, 0.3463, 0.0755]
sd_o = float(np.std(orig, ddof=1))
sd_i = float(np.std(impr, ddof=1))
print(f'原版样本 SD = {sd_o:.3f}；改进版样本 SD = {sd_i:.3f}')

REPL = [
    ('(mean 0.131, SD 0.197)', f'(mean 0.131, sample SD {sd_o:.3f})'),
    ('(mean 0.131, SD 0.197)', f'(mean 0.131, sample SD {sd_o:.3f})'),
    ('(SD 0.197 on the side-chain-centroid definition)',
     f'(sample SD {sd_o:.3f} on the side-chain-centroid definition)'),
    ('SD 0.197) and a rescaled coordinate-update variant did not stabilise it (SD 0.216)',
     f'sample SD {sd_o:.3f}) and a rescaled coordinate-update variant did not stabilise it '
     f'(sample SD {sd_i:.3f})'),
    ('(SD 0.197 across three seeds)', f'(sample SD {sd_o:.3f} across three seeds)'),
]


def main():
    doc = Document(DOC)
    n = 0
    for p in doc.paragraphs:
        if not p.runs:
            continue
        full = ''.join(r.text for r in p.runs)
        new = full
        for a, b in REPL:
            if a in new:
                new = new.replace(a, b)
                n += 1
        if new != full:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ''
    doc.save(DOC)
    print(f'替换 {n} 处')
    d2 = Document(DOC)
    txt = '\n'.join(p.text for p in d2.paragraphs)
    print(f'残留 SD 0.197: {txt.count("SD 0.197")} | 残留 SD 0.216: {txt.count("SD 0.216")}')
    print(f'新 SD 出现: {txt.count(f"{sd_o:.3f}")} 次（原版）, {txt.count(f"{sd_i:.3f}")} 次（改进版）')


if __name__ == '__main__':
    main()
