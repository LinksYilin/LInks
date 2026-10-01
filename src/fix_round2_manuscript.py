# -*- coding: utf-8 -*-
"""fix_round2_manuscript.py — 修正第二轮验证发现的论文缺陷"""
from docx import Document

DOC = r'D:\GED_mutation\manuscript_routeA.docx'
A = '\u03b1'
B = '\u03b2'
RHO = '\u03c1'
DEL = '\u0394'
MIN = '\u2212'

RULES = [
    # A-1: Figure 3 图注"四个定义" → 如实说明某两档只有三个
    ('a, S669 Pearson r for five encoders spanning 5.8 k to 509 k parameters, each trained on the '
     'audited training set under four graph definitions (C' + A + ', C' + B + ', side-chain '
     'centroid, all-atom).',
     'a, S669 Pearson r for five encoders spanning 5.8 k to 509 k parameters, each trained on the '
     'audited training set. The three lower-capacity encoders were trained under all four graph '
     'definitions (C' + A + ', C' + B + ', side-chain centroid, all-atom); the two largest were '
     'trained under three, because all-atom graphs for those architectures were not run.'),

    # A-1: Table 1 脚注 †
    ('Encoder results on S669 are reported as the range of three-seed mean-of-per-seed '
     'correlations across the C' + A + ', C' + B + ', side-chain-centroid and all-atom definitions;',
     'Encoder results on S669 are reported as the range of three-seed mean-of-per-seed '
     'correlations across the graph definitions trained for each encoder: C' + A + ', C' + B +
     ', side-chain centroid and all-atom for the three lower-capacity encoders, and C' + A + ', C' +
     B + ' and side-chain centroid for the 460 k and 509 k encoders;'),

    # A-3 / D-4: Table 1 脚注 +0.010 → +0.009
    ('on S669 the increment is excluded below ' + MIN + '0.036 and above +0.010;',
     'on S669 the increment is excluded below ' + MIN + '0.036 and above +0.009;'),

    # B-1: 五个 → 六个
    ('Five boundaries define the scope of these conclusions.',
     'Six boundaries define the scope of these conclusions.'),

    # B-2: ssym ridge 0.306 → 0.305
    ('r = 0.306 on ssym', 'r = 0.305 on ssym'),
    ('0.306 on ssym', '0.305 on ssym'),

    # B-3: 0.36-0.41 → 0.36-0.40
    ('r = 0.36-0.41 for local pooling', 'r = 0.36-0.40 for local pooling'),
    ('0.36\u20130.41 for local pooling', '0.36\u20130.40 for local pooling'),

    # B-4: 0.062 → 0.063
    ('moves by up to 0.062 on ssym', 'moves by up to 0.063 on ssym'),

    # B-5: +0.01 → +0.009
    ('no positive increment above +0.01 on S669', 'no positive increment above +0.009 on S669'),

    # B-6: "at matched capacity" 无支撑 → 改为如实描述
    ('encoders with explicit edge features did not outperform those without them at matched '
     'capacity.',
     'the one encoder with explicit edge features and the one without them that we compared '
     'directly (50 k versus 5.9 k parameters) did not differ significantly, although those two '
     'differ in capacity as well as in edge handling.'),

    # B-7: VIF 声明需说明设计
    ('all variance inflation factors were below 2.4 times.',
     'variance inflation factors stayed below 2.4 when the physicochemical features were entered '
     'together with a single contact count; entering the three counts jointly is not identifiable, '
     'because the total edit count is their exact sum.'),

    # B-9: 破损句子修正（大写 A 开头）
    ('not reproduced baselines in this comparison, and A representation effect remains possible '
     'in encoder families outside the five tested here.',
     'not reproduced baselines in this comparison. A representation effect remains possible in '
     'encoder families outside the five tested here.'),

    # E-4: Methods 中未被使用的 BLOSUM62 特征 → 删除并说明
    ('Node-edit feature: BLOSUM62 substitution score [6] between the wild-type and mutant amino '
     'acids placed at the mutated residue node.',
     'Node-edit feature: the substituted residue identity and the BLOSUM62 substitution score [6] '
     'are available to the model through the one-hot mutant column and the physicochemical '
     'property columns of the mutated node. The ladder models use no separate BLOSUM62 input '
     'column; that feature belongs to earlier iterations and to the ablation in Section 4.8.'),
    ('a fixed BLOSUM62 node-edit feature is appended as an additional column.',
     'the mutant residue identity at the mutated position is encoded in the node features.'),
]


def main():
    d = Document(DOC)
    n = 0
    hits = []
    for p in d.paragraphs:
        if not p.runs:
            continue
        full = ''.join(r.text for r in p.runs)
        new = full
        for a, b in RULES:
            if a in new:
                new = new.replace(a, b)
                hits.append(a[:56])
        if new != full:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ''
            n += 1
    d.save(DOC)
    print(f'修改 {n} 段; 命中 {len(hits)} 条:')
    for h in hits:
        print(f'  ✓ {h}')

    d2 = Document(DOC)
    txt = '\n'.join(p.text for p in d2.paragraphs) + '\n' + '\n'.join(
        c.text for t in d2.tables for r in t.rows for c in r.cells)
    print()
    for pat, name in [('Five boundaries', 'Five boundaries'), ('0.306 on ssym', '0.306'),
                      ('at matched capacity', 'matched capacity'),
                      ('all-atom for the three lower', '新 Figure3 说明'),
                      ('Six boundaries', 'Six boundaries')]:
        print(f'  {name}: {txt.count(pat)}')


if __name__ == '__main__':
    main()
