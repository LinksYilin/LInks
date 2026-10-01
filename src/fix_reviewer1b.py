# -*- coding: utf-8 -*-
"""
fix_reviewer1b.py — R1-M3（分类型接触的准确表述与多重比较）+ R1-M6（0.402 的来源）
"""
from docx import Document

DOC = r'D:\GED_mutation\manuscript_routeA.docx'
D = '\u0394'
M = '\u2212'

FIX = [
    # --- [138] Results 4.5 ---
    ('Stratifying by contact type did produce one significant quantity: the number of broken '
     'hydrophobic contacts correlated with ' + D + D + 'G at r = 0.135 (95% CI 0.044 to 0.201), '
     'the only contact-derived feature to exclude zero. Broken electrostatic (r = 0.047), other '
     'broken (r = 0.027), formed hydrophobic (r = 0.096), formed electrostatic (r = ' + M + '0.046) '
     'and other formed (r = 0.063) contacts were all non-significant.',
     'Stratifying by contact type produced two nominally significant quantities among the six tests: '
     'the number of broken hydrophobic contacts (r = 0.135, 95% CI 0.044 to 0.201, nominal '
     'P = 0.0024) and the number of formed hydrophobic contacts (r = 0.096, nominal P = 0.032). '
     'Broken electrostatic (r = 0.047), other broken (r = 0.027), formed electrostatic '
     '(r = ' + M + '0.046) and other formed (r = 0.063) contacts were not significant. Only the '
     'broken hydrophobic count survives a Bonferroni correction across the six tests (threshold '
     '0.0083), and we therefore treat it as the single quantity that withstands multiplicity '
     'control rather than as a feature selected from six.'),

    # --- [157] Results 4.8 ---
    ('Stratifying by contact type isolated one significant quantity, the count of broken '
     'hydrophobic contacts (r = 0.135, 95% CI 0.044 to 0.201), which nonetheless remained '
     'redundant with the physicochemical features once both were entered together.',
     'Stratifying by contact type isolated two nominally significant counts, of broken '
     '(r = 0.135, nominal P = 0.0024) and formed (r = 0.096, nominal P = 0.032) hydrophobic '
     'contacts; of these only the broken count survives a Bonferroni correction across six tests. '
     'The broken count nonetheless remained redundant with the physicochemical features once both '
     'were entered together.'),

    # --- [165] Figure 5 图注 ---
    ('Only the count of broken hydrophobic contacts excludes zero (r = 0.135, 95% CI 0.044 to '
     '0.201, highlighted); the remaining five type-specific counts are non-significant.',
     'Two counts are nominally significant, of broken (r = 0.135, 95% CI 0.044 to 0.201, '
     'highlighted) and formed (r = 0.096) hydrophobic contacts; only the broken count survives a '
     'Bonferroni correction across the six tests (threshold 0.0083).'),

    # --- [186] Discussion 5.2 ---
    ('adds no incremental information over those features (cross-validated ' + D + 'r = -0.009, '
     '95% CI -0.047 to +0.037)',
     'adds no increment above about +0.04 over those features (cross-validated ' + D + 'r = '
     + M + '0.009, 95% CI ' + M + '0.047 to +0.037)'),

    # --- R1-M6：0.402 是两种子均值，应说明 ---
    ('The highest single value among graph encoders was r = 0.402, attained by the 50 k-parameter '
     'GINE on the C' + '\u03b2' + ' definition, which exceeds the physicochemical baseline.',
     'The highest value among graph encoders was r = 0.402 (two-seed mean; individual seeds 0.418 '
     'and 0.387), attained by the 50 k-parameter GINE on the C' + '\u03b2' + ' definition, which '
     'exceeds the physicochemical baseline.'),

    ('The 50 k-parameter GINE encoder attained the highest single S669 value among graph encoders '
     '(r = 0.402 on C' + '\u03b2' + '), above the physicochemical baseline.',
     'The 50 k-parameter GINE encoder attained the highest S669 value among graph encoders '
     '(r = 0.402, two-seed mean, on C' + '\u03b2' + '), above the physicochemical baseline.'),

    ('The 50 k-parameter GINE encoder reached r = 0.402 on the C' + '\u03b2' + ' definition, the '
     'highest value among the graph encoders and above the physicochemical ridge baseline '
     '(r = 0.390)',
     'The 50 k-parameter GINE encoder reached r = 0.402 (two-seed mean) on the C' + '\u03b2' +
     ' definition, the highest value among the graph encoders and above the physicochemical ridge '
     'baseline (r = 0.390)'),
]


def main():
    d = Document(DOC)
    n = 0
    for p in d.paragraphs:
        if not p.runs:
            continue
        full = ''.join(r.text for r in p.runs)
        new = full
        for a, b in FIX:
            if a in new:
                new = new.replace(a, b)
                n += 1
        if new != full:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ''
    d.save(DOC)
    print(f'修正 {n} 处')
    d2 = Document(DOC)
    for p in d2.paragraphs:
        if 'two nominally significant quantities among the six tests' in p.text:
            print('  ✅ R1-M3 (4.5) 已修')
        if 'two-seed mean' in p.text:
            print('  ✅ R1-M6 已修')
            break


if __name__ == '__main__':
    main()
